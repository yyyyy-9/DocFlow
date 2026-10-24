"""Real loopback API: authentication, first draft, classification and feedback."""
from pathlib import Path
from urllib.request import Request, build_opener, ProxyHandler
from urllib.error import HTTPError
from urllib.parse import urlsplit, parse_qs
import json, subprocess, sys, tempfile, time, unittest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]


class UIWorkflow(unittest.TestCase):
    def test_api_three_drafts_and_source_preservation(self):
        with tempfile.TemporaryDirectory(prefix='docflow-ui-') as name:
            home = Path(name); data = home/'state'; source = home/'source'; source.mkdir()
            originals = {}
            for n in (1, 2):
                path = source/f'{n:02d}-report.docx'; doc = Document()
                doc.add_paragraph(f'试验报告{n}', style='Title')
                doc.add_paragraph('目的', style='Heading 1')
                doc.add_paragraph(f'独立测量{n}：2.30 V 与 1.5 mm 必须保持。')
                doc.save(path); originals[path] = path.read_bytes()
            with (home/'server.log').open('wb') as log:
                process = subprocess.Popen([sys.executable, '-B', str(ROOT/'WPS_Forms_UI/ui_tool.py'), 'serve', '--data', str(data), '--no-browser'], stdout=log, stderr=subprocess.STDOUT)
                try:
                    deadline = time.monotonic() + 30
                    while not (data/'server.json').is_file():
                        self.assertIsNone(process.poll(), 'UI server exited during startup')
                        self.assertLess(time.monotonic(), deadline, 'UI startup timeout')
                        time.sleep(.05)
                    server = json.loads((data/'server.json').read_text('utf8'))
                    uri = urlsplit(server['url']); base = f'{uri.scheme}://{uri.netloc}'
                    token = parse_qs(uri.fragment)['token'][0]

                    opener = build_opener(ProxyHandler({}))
                    def request(path, payload=None, authorized=True, origin=None):
                        headers = {'Content-Type': 'application/json'}
                        if authorized: headers['X-WPS-UI-Token'] = token
                        if origin: headers['Origin'] = origin
                        req = Request(base + path, data=None if payload is None else json.dumps(payload).encode(), headers=headers)
                        with opener.open(req, timeout=30) as response: return json.load(response)

                    for kwargs in ({'authorized': False}, {'origin': 'https://example.com'}):
                        with self.assertRaises(HTTPError) as caught: request('/api/state', **kwargs)
                        self.assertEqual(caught.exception.code, 403)
                    self.assertEqual(request('/api/state')['forms_version'], '0.1.6')

                    def submit(operation, **payload):
                        request('/api/submit', dict(operation=operation, origin='model', **payload))
                        deadline = time.monotonic() + 90
                        while True:
                            state = request('/api/state'); job = state['job']
                            if job['status'] not in ('RUNNING', 'DELIVERED', 'DELIVERED_WITH_ISSUES'): break
                            if job['status'].startswith('DELIVERED') and state['current'] == job['output']: break
                            self.assertLess(time.monotonic(), deadline, f'{operation} timeout')
                            time.sleep(.1)
                        self.assertEqual(job['status'], 'DELIVERED', job)
                        self.assertTrue(Path(state['master']).is_file())
                        return state

                    first = submit('scan', root=str(source), title='公开测试', pages=False, doc_conversion=False)
                    first_path = Path(first['master']); first_bytes = first_path.read_bytes()
                    form = request('/api/form?kind=02')
                    synced = request('/api/sync', dict(kind='02', version=form['version'], changes={'F9': '归类A', 'H9': '展示甲'}, origin='model'))
                    with self.assertRaises(HTTPError) as caught:
                        request('/api/submit', dict(operation='classify', version=form['version'], origin='model'))
                    self.assertEqual(caught.exception.code, 409)
                    second = submit('classify', version=synced['version'])
                    self.assertIn('展示甲', '\n'.join(p.text for p in Document(second['master']).paragraphs))
                    form = request('/api/form?kind=03')
                    row = next(item['row'] for item in form['rows'] if not item['values']['A'])
                    values = {'B': 'master', 'C': '独立测量1：2.30 V 与 1.5 mm 必须保持。', 'D': '字号', 'E': '14', 'G': '仅此处', 'H': '提交修改'}
                    synced = request('/api/sync', dict(kind='03', version=form['version'], changes={col+str(row): value for col,value in values.items()}, origin='model'))
                    third = submit('feedback', version=synced['version'])
                    para = next(p for p in Document(third['master']).paragraphs if p.text == values['C'])
                    self.assertEqual(para.runs[0].font.size.pt, 14)
                    self.assertEqual(len(third['history']), 3)
                    self.assertEqual(first_path.read_bytes(), first_bytes)
                    for path, content in originals.items(): self.assertEqual(path.read_bytes(), content)
                    request('/api/shutdown', {})
                    self.assertEqual(process.wait(timeout=15), 0)
                finally:
                    if process.poll() is None:
                        process.terminate()
                        try: process.wait(timeout=10)
                        except subprocess.TimeoutExpired: process.kill(); process.wait(timeout=5)


if __name__ == '__main__': unittest.main()
