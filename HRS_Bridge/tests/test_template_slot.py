"""Consume WPS's declared body slot, preserving native front matter and fields."""
from pathlib import Path
import os,json,tempfile,zipfile,unittest,shutil
from docx import Document
from docx.oxml.ns import qn
from lxml import etree as E
from hrs_bridge.runtime import sha,BridgeError
from hrs_bridge.docx_export import convert
from hrs_bridge.docx_import import convert as read_back
from hrs_bridge.ooxml import N,bookmark_regions,provenance_path
from test_html import runtime

ROOT=Path(__file__).resolve().parents[1]

class TemplateSlot(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.tmp=tempfile.TemporaryDirectory(prefix='bridge-body-slot-');cls.folder=Path(cls.tmp.name)
  cls.rt=runtime();cls.source=cls.folder/'project';(cls.source/'assets').mkdir(parents=True)
  cls.raw=('# 槽位回归报告\n\n## 试验目的\n\n正文说明。\n\n1. 列表甲\n2. 列表乙\n\n### 细节\n\n'
   '```report-table\n{"schema":"report-table-v1","id":"tbl-one","caption":"参数","rows":[[{"text":"值","rowspan":1,"colspan":1}]]}\n```\n\n'
   '参见 @tbl:tbl-one。\n\n## 结论\n\n1. 新列表\n').encode('utf8')
  (cls.source/'report.md').write_bytes(cls.raw)
  (cls.source/'meta.json').write_text(json.dumps({'render_profile':'free','render_options':{'show_object_numbers':False}}),encoding='utf8')
  cls.template=ROOT/'fixtures/cover_template_v0311.docx'
  cls.template_raw=cls.template.read_bytes();cls.cover=cls.folder/'cover.docx';cls.default=cls.folder/'default.docx'
  cls.cover_receipt=convert(cls.rt,cls.source,cls.cover,template=cls.template)
  cls.default_receipt=convert(cls.rt,cls.source,cls.default)
  cls.doc=Document(cls.cover)
 @classmethod
 def tearDownClass(cls):
  assert cls.template.read_bytes()==cls.template_raw
  evidence=os.environ.get('HRS_RC2_EVIDENCE')
  if evidence:
   destination=Path(evidence);destination.mkdir(parents=True,exist_ok=True)
   for file in cls.folder.glob('*.docx'):shutil.copyfile(file,destination/file.name)
   for file in cls.folder.glob('*.hrs-bridge.json'):shutil.copyfile(file,destination/file.name)
   (destination/'receipts.json').write_text(json.dumps({'default':cls.default_receipt,'cover':cls.cover_receipt},ensure_ascii=False,indent=2),encoding='utf8')
  cls.tmp.cleanup()
 def test_front_body_and_sections_are_unchanged(self):
  from wps_pipeline.resolve_styles import inspect_template
  original=Document(self.template);slot=inspect_template(self.template)['body_slot']['index']
  # Namespace prefixes are serialization details; subtree content/format stays equal.
  canonical=lambda node:E.tostring(E.fromstring(E.tostring(node)),method='c14n')
  self.assertEqual([canonical(n) for n in list(original._element.body)[:slot]],
                   [canonical(n) for n in list(self.doc._element.body)[:slot]])
  self.assertEqual(canonical(original._element.body[-1]),canonical(self.doc._element.body[-1]))
  self.assertFalse(self.doc._element.body.xpath('.//w:sdt[w:sdtPr/w:tag[@w:val="wps-report-body-v1"]]'))
 def test_headers_footers_and_opaque_parts_are_byte_exact(self):
  mutable={'word/document.xml','word/styles.xml','word/numbering.xml','word/settings.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
  with zipfile.ZipFile(self.template) as original,zipfile.ZipFile(self.cover) as output:
   for name in original.namelist():
    if name not in mutable:self.assertEqual(original.read(name),output.read(name),name)
 def test_live_heading_list_caption_and_reference_structures(self):
  title=next(p for p in self.doc.paragraphs if p.text=='槽位回归报告');self.assertEqual(title.style.name,'Title')
  heading=next(p for p in self.doc.paragraphs if p.text=='试验目的');self.assertEqual(heading.style.name,'Heading 1')
  instructions=self.doc._element.xpath('.//w:fldSimple/@w:instr')+self.doc._element.xpath('.//w:instrText/text()')
  self.assertTrue(any('SEQ' in i for i in instructions));self.assertTrue(any('REF' in i for i in instructions));self.assertTrue(any('TOC' in i for i in instructions))
  list_ids=[p._p.find('.//'+qn('w:numId')).get(qn('w:val')) for p in self.doc.paragraphs if p.text in ('列表甲','列表乙','新列表')]
  self.assertEqual(list_ids[0],list_ids[1]);self.assertNotEqual(list_ids[1],list_ids[2])
  self.assertEqual(self.doc.settings.element.find(qn('w:updateFields')).get(qn('w:val')),'true')
 def test_cover_does_not_enter_roundtrip_markdown(self):
  out=self.folder/'back';read_back(self.rt,self.cover,out)
  self.assertEqual((out/'report.md').read_bytes(),self.raw)
  from render_profile import normalize_options
  self.assertEqual(json.loads((out/'meta.json').read_text())['render_options'],normalize_options({'show_object_numbers':False}))
 def test_missing_body_boundary_and_untracked_tail_fail(self):
  meta=json.loads(provenance_path(self.cover).read_text());name=meta['body_scope']['bookmark']
  doc=Document(self.cover);node=next(n for n in doc._element.iter(qn('w:bookmarkStart')) if n.get(qn('w:name'))==name);node.getparent().remove(node)
  with self.assertRaises(BridgeError):bookmark_regions(doc,name)
  doc=Document(self.cover);doc.add_paragraph('未映射的新正文')
  with self.assertRaisesRegex(BridgeError,'UNTRACKED_CONTENT'):bookmark_regions(doc,name)
 def test_body_boundary_moved_into_paragraphs_still_slices(self):
  meta=json.loads(provenance_path(self.cover).read_text());name=meta['body_scope']['bookmark'];doc=Document(self.cover);body=doc._element.body
  start=next(n for n in body if n.tag==qn('w:bookmarkStart') and n.get(qn('w:name'))==name);sid=start.get(qn('w:id'))
  end=next(n for n in body if n.tag==qn('w:bookmarkEnd') and n.get(qn('w:id'))==sid)
  # WPS may put range markers inside the first/last paragraph during save.
  first=next(n for n in list(body)[list(body).index(start)+1:] if n.tag==qn('w:p'))
  last=next(n for n in reversed(list(body)[:list(body).index(end)]) if n.tag==qn('w:p'))
  first_block=next(n for n in list(body)[list(body).index(start)+1:] if n.tag==qn('w:bookmarkStart'))
  last_block=next(n for n in reversed(list(body)[:list(body).index(end)]) if n.tag==qn('w:bookmarkEnd'))
  first.insert(1,start);first.insert(2,first_block);last.append(last_block);last.append(end)
  self.assertEqual(len(bookmark_regions(doc,name)),len(meta['blocks']))
 def test_shared_engineering_policy_rejects_bad_titles(self):
  source=self.folder/'bad';(source/'assets').mkdir(parents=True)
  for name,text,code in [('double','# 一\n\n# 二\n','MULTIPLE_H1'),('gap','# 一\n\n#### 四\n','HEADING_LEVEL_GAP'),('prefix','# 一\n\n## 1 二\n','HEADING_SOURCE_PREFIX')]:
   (source/'report.md').write_bytes(text.encode())
   with self.assertRaisesRegex(ValueError,code):convert(self.rt,source,self.folder/(name+'.docx'))
   self.assertFalse((self.folder/(name+'.docx')).exists())
 def test_wps_coincident_scope_markers_can_swap_without_losing_body(self):
  meta=json.loads(provenance_path(self.cover).read_text());name=meta['body_scope']['bookmark'];doc=Document(self.cover);body=doc._element.body
  starts=list(body.iter(qn('w:bookmarkStart')))
  scope=next(n for n in starts if n.get(qn('w:name'))==name);sid=scope.get(qn('w:id'))
  first=next(n for n in starts if n.get(qn('w:name'))==meta['blocks'][0]['bookmark'])
  end=next(n for n in body.iter(qn('w:bookmarkEnd')) if n.get(qn('w:id'))==sid)
  last_name=meta['blocks'][-1]['bookmark'];last=next(n for n in starts if n.get(qn('w:name'))==last_name)
  last_end=next(n for n in body.iter(qn('w:bookmarkEnd')) if n.get(qn('w:id'))==last.get(qn('w:id')))
  first.addnext(scope);last_end.addprevious(end)
  self.assertEqual(len(bookmark_regions(doc,name)),len(meta['blocks']))
  # The tolerated swap is zero-width: any intervening new content is rejected.
  gap=E.Element(qn('w:r'));E.SubElement(gap,qn('w:t')).text='outside scope';first.addnext(gap)
  with self.assertRaisesRegex(BridgeError,'BRIDGE_BODY_SCOPE_ORDER'):bookmark_regions(doc,name)

if __name__=='__main__':unittest.main()
