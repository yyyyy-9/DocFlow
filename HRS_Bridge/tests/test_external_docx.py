"""INT-01: harmless Word markers and an explicitly approved external image group."""
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from PIL import Image

from hrs_bridge.docx_import import convert
from hrs_bridge.runtime import BridgeError, sha
from test_html import runtime, ROOT


class ExternalDOCX(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='bridge-external-')
        self.folder=Path(self.temp.name)
        self.rt=runtime()

    def tearDown(self):
        self.temp.cleanup()

    def test_empty_editor_markers_do_not_hide_real_paragraph_or_cell_text(self):
        """WPS spell-check and rendered-page markers are metadata; adjacent content must survive."""
        doc=Document(ROOT/'fixtures/target_template.docx')
        p=doc.add_paragraph('正文甲')
        p._p.insert(1,OxmlElement('w:proofErr'))
        p.add_run()._r.append(OxmlElement('w:lastRenderedPageBreak'))
        p.add_run('正文乙')
        table=doc.add_table(rows=1,cols=1)
        table.cell(0,0).text='单元格内容'
        table.cell(0,0).paragraphs[0]._p.insert(1,OxmlElement('w:proofErr'))
        source=self.folder/'markers.docx';doc.save(source)
        result=convert(self.rt,source,self.folder/'back')
        md=(self.folder/'back/report.md').read_text('utf8')
        self.assertIn('正文甲正文乙',md)
        self.assertIn('单元格内容',md)
        self.assertEqual(sum(d['code']=='WORD_EDITOR_MARKER_IGNORED' for d in result['diagnostics']),3)

    def test_content_inside_an_editor_marker_is_not_discarded(self):
        """Only empty markers may be ignored; unusual content must still stop conversion."""
        doc=Document(ROOT/'fixtures/target_template.docx')
        marker=OxmlElement('w:proofErr');marker.text='不可丢失'
        doc.add_paragraph('正文')._p.append(marker)
        source=self.folder/'nonempty.docx';doc.save(source)
        with self.assertRaisesRegex(BridgeError,'UNSUPPORTED_NATIVE_CONTENT'):
            convert(self.rt,source,self.folder/'rejected')
        self.assertFalse((self.folder/'rejected').exists())

    def test_confirmed_body_number_remains_literal_body_text(self):
        """A source-bound body decision preserves typed numbers regardless of paragraph style."""
        doc=Document(ROOT/'fixtures/target_template.docx')
        doc.add_paragraph('2. 测试',style='Normal')
        source=self.folder/'literal.docx';doc.save(source)
        from wps_pipeline.scan_docx import Scanner
        scan=Scanner(source,self.folder/'literal-scan').run()
        block=next(b for b in scan['blocks'] if b.get('plain_text')=='2. 测试')
        plan={'schema':'hrs-word-structure-map-v1','source_sha256':sha(source.read_bytes()),'roles':{
            block['block_id']:{'role':'body','source_sha256':sha(source.read_bytes()),'reason':'User confirmed literal body text'}}}
        decision=self.folder/'body.json';decision.write_text(json.dumps(plan),encoding='utf8')
        convert(self.rt,source,self.folder/'literal-back',structure_map=decision)
        raw=(self.folder/'literal-back/report.md').read_text('utf8')
        self.assertIn('2\\. 测试',raw)

    def make_images(self):
        doc=Document(ROOT/'fixtures/target_template.docx')
        image=self.folder/'pic.png';Image.new('RGB',(32,24),'navy').save(image)
        for row in range(2):
            p=doc.add_paragraph()
            for col in range(2):p.add_run().add_picture(str(image),width=Inches(.7))
        doc.add_paragraph('测试1    测试2')
        source=self.folder/'four.docx';doc.save(source)
        from wps_pipeline.scan_docx import Scanner
        scan=Scanner(source,self.folder/'scan').run()
        blocks=[b['block_id'] for b in scan['blocks'] if b['kind']=='paragraph' and any(s['kind']=='drawing' for s in b['segments'])]
        plan={'schema':'hrs-word-structure-map-v1','source_sha256':sha(source.read_bytes()),
            'groups':[{'blocks':blocks,'columns':2,'rows':2,'showSubfigureLabels':False,'caption':''}]}
        path=self.folder/'structure.json';path.write_text(json.dumps(plan),encoding='utf8')
        return source,path,plan

    def test_explicit_source_bound_group_keeps_four_images_and_following_text(self):
        """User-confirmed adjacent picture paragraphs become one 2x2 group; no text is consumed."""
        source,path,plan=self.make_images()
        result=convert(self.rt,source,self.folder/'group',structure_map=path)
        raw=(self.folder/'group/report.md').read_bytes()
        ir=self.rt.analyze(raw,self.folder/'group/assets')
        groups=[b for b in ir['blocks'] if b['type']=='image-group']
        self.assertEqual(len(groups),1)
        self.assertEqual((groups[0]['data']['columns'],groups[0]['data']['rows']),(2,2))
        self.assertEqual(len(groups[0]['data']['images']),4)
        self.assertFalse(groups[0]['data']['showSubfigureLabels'])
        self.assertIn('测试1    测试2',raw.decode('utf8'))
        self.assertEqual(result['mode'],'external-structured')

    def test_picture_paragraphs_are_separate_rows_without_cross_row_decision(self):
        """Automatic row grouping keeps both rows, without inferring a shared cross-row object."""
        source,path,plan=self.make_images()
        convert(self.rt,source,self.folder/'rows')
        ir=self.rt.analyze((self.folder/'rows/report.md').read_bytes(),self.folder/'rows/assets')
        groups=[b for b in ir['blocks'] if b['type']=='image-group']
        self.assertEqual(len(groups),2)
        self.assertTrue(all((b['data']['rows'],b['data']['columns'])==(1,2) for b in groups))
        self.assertIn('测试1    测试2',(self.folder/'rows/report.md').read_text('utf8'))

    def test_group_mapping_cannot_be_applied_to_another_source(self):
        """A local decision becomes invalid when the bound DOCX bytes change."""
        source,path,plan=self.make_images()
        plan['source_sha256']='0'*64
        path.write_text(json.dumps(plan),encoding='utf8')
        with self.assertRaisesRegex(BridgeError,'STRUCTURE_MAP_SOURCE_MISMATCH'):
            convert(self.rt,source,self.folder/'wrong',structure_map=path)

    def test_group_decision_cannot_consume_adjacent_body_text(self):
        """A declared group cannot swallow a following paragraph, even if its ID is supplied."""
        source,path,plan=self.make_images()
        scan=json.loads((self.folder/'scan/scan.json').read_text('utf8'))
        body=next(b for b in scan['blocks'] if b.get('plain_text')=='测试1    测试2')
        plan['groups'][0]['blocks'].append(body['block_id'])
        path.write_text(json.dumps(plan),encoding='utf8')
        with self.assertRaisesRegex(BridgeError,'consume non-image content'):
            convert(self.rt,source,self.folder/'no-text-loss',structure_map=path)


if __name__=='__main__':unittest.main()
