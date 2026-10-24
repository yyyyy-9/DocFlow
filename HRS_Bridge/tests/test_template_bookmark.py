"""TPL-03: new plain-body template in both production dispatch paths."""
from pathlib import Path
import json,tempfile,unittest,zipfile
from docx import Document
from lxml import etree as E
from test_html import runtime
from hrs_bridge.runtime import sha
from hrs_bridge.docx_export import convert
from hrs_bridge.docx_import import convert as restore
from hrs_bridge.workflow import run
ROOT=Path(__file__).resolve().parents[1]
N={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

class BookmarkTemplate(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.tmp=tempfile.TemporaryDirectory();cls.folder=Path(cls.tmp.name);cls.rt=runtime()
  from wps_pipeline.resolve_styles import inspect_template
  from wps_pipeline.template_prepare import prepare_template
  old=ROOT/'fixtures/cover_template_v0311.docx';cls.template=cls.folder/'plain.docx'
  slot=inspect_template(old)['body_slot']
  if slot['kind']=='bookmark':cls.template.write_bytes(old.read_bytes())
  else:prepare_template(old,cls.template,slot['index'],sha(old.read_bytes()))
  cls.project=cls.folder/'project';(cls.project/'assets').mkdir(parents=True)
  cls.raw='# 实际报告题名\n\n## 方法\n\n正文内容。\n\n## 结果\n\n结果说明。\n'.encode('utf8')
  (cls.project/'report.md').write_bytes(cls.raw)
 @classmethod
 def tearDownClass(cls):cls.tmp.cleanup()
 def assert_plain_output(self,path):
  from wps_pipeline.resolve_styles import inspect_template
  with zipfile.ZipFile(self.template) as before,zipfile.ZipFile(path) as after:
   template=E.fromstring(before.read('word/document.xml'));out=E.fromstring(after.read('word/document.xml'))
   prefix=inspect_template(self.template)['body_slot']['index']
   canonical=lambda n:E.tostring(n,method='c14n')
   self.assertEqual([canonical(n) for n in template.find('w:body',N)[:prefix]],[canonical(n) for n in out.find('w:body',N)[:prefix]])
   self.assertFalse(out.xpath('.//w:bookmarkStart[@w:name="WPS_ReportBody"]',namespaces=N))
   self.assertFalse(out.xpath('.//w:sdt[w:sdtPr/w:tag[@w:val="wps-report-body-v1"]]',namespaces=N))
   for name in before.namelist():
    if name.startswith(('word/header','word/footer')):self.assertEqual(before.read(name),after.read(name))
 def test_hrs_word_roundtrip_keeps_body_identity_and_excludes_front_pages(self):
  path=self.folder/'hrs.docx';result=convert(self.rt,self.project,path,template=self.template)
  self.assertEqual(result['template']['body_slot']['kind'],'bookmark');self.assert_plain_output(path)
  restored=self.folder/'back';restore(self.rt,path,restored)
  self.assertEqual((restored/'report.md').read_bytes(),self.raw)
 def test_word_native_task_uses_new_bookmark_and_finishes_independent_checks(self):
  source=self.folder/'source.docx';doc=Document();doc.add_paragraph('1 试验方法',style='Heading 1');doc.add_paragraph('测量值 1.2毫米，原文保留。');doc.save(source)
  job=self.folder/'job.json';job.write_text(json.dumps({'schema':'hrs-bridge-job-v2','output':'native',
   'template':str(self.template),'tasks':[{'id':'word','source':{'kind':'docx','path':str(source)},'outputs':['docx']}]}),encoding='utf8')
  result=run(self.rt,job)
  if result['status']!='COMPLETE':
   failures={p.name:json.loads(p.read_text('utf8')).get('errors') for p in (self.folder/'native').rglob('*.content_check.json')}
   self.fail({'status':result['status'],'content_errors':failures})
  task=result['tasks'][0];self.assertEqual(task['native_status']['receipt_integrity'],'PASS')
  for deliverable in task['deliverables']:self.assert_plain_output(self.folder/'native'/deliverable['path'])

if __name__=='__main__':unittest.main()
