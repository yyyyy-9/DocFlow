"""B4: current native content wins; provenance restores only unchanged source."""
from pathlib import Path
import tempfile
import unittest
from docx import Document
from hrs_bridge.browser import recover
from hrs_bridge.docx_export import convert as to_docx
from hrs_bridge.ooxml import N, W, read_provenance, bookmark_regions
from hrs_bridge.runtime import BridgeError
from test_html import ROOT, runtime


class DOCXImport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from hrs_bridge.docx_import import convert
        cls.convert = staticmethod(convert)
        cls.temp = tempfile.TemporaryDirectory(prefix='bridge-import-test-')
        cls.folder = Path(cls.temp.name)
        cls.rt = runtime()
        cls.source = cls.folder / 'source'
        recover(cls.rt, ROOT / 'fixtures/hrs_sample.html', cls.source)
        cls.docx = cls.folder / 'original.docx'
        cls.receipt = to_docx(cls.rt, cls.source, cls.docx)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_unchanged_native_document_restores_source_and_assets_exactly(self):
        """B4: shared resources and source kind/ID/numbering survive a same-source round trip."""
        out = self.folder / 'exact'
        result = self.convert(self.rt, self.docx, out)
        self.assertTrue(result['ok'])
        self.assertEqual((out / 'report.md').read_bytes(), (self.source / 'report.md').read_bytes())
        self.assertEqual({p.name:p.read_bytes() for p in (out/'assets').iterdir()},
                         {p.name:p.read_bytes() for p in (self.source/'assets').iterdir()})

    def test_edited_table_cell_is_read_from_word(self):
        """B4: changing a native cell must retain tbl ID and must never restore its stale text."""
        d = Document(self.docx)
        d.tables[2].cell(1, 1).text = 'WORD_EDIT_123'
        edited = self.folder / 'edited.docx'
        d.save(edited)
        out = self.folder / 'edited-back'
        result = self.convert(self.rt, edited, out, provenance=Path(self.receipt['provenance']))
        ir = self.rt.analyze((out/'report.md').read_bytes(), out/'assets')
        b = next(b for b in ir['blocks'] if b.get('id') == 'tbl-deebf0cf6ded')
        # Cell text is Markdown source. Compare its actual HRS inline rendering,
        # allowing the source serializer to protect literal punctuation.
        from hrs_bridge.browser import call
        tokens = call(self.rt, 'tokens', markdown=b['data']['rows'][1][1]['text'])
        self.assertEqual(''.join(t['text'] for t in tokens[0]['tokens']), 'WORD_EDIT_123')
        self.assertEqual(len(result['changed_blocks']), 1)

    def test_missing_provenance_does_not_guess_group_identity(self):
        """B4: bridge bookmarks without their mapping are a failure, not an external-table fallback."""
        d = Document(self.docx)
        missing = self.folder / 'missing-map.docx'
        d.save(missing)
        with self.assertRaisesRegex(BridgeError, 'PROVENANCE_REQUIRED'):
            self.convert(self.rt, missing, self.folder / 'no-map')
        self.assertFalse((self.folder/'no-map').exists())

    def test_native_equation_edit_keeps_identity_and_reads_current_omml(self):
        """B4: current OMML replaces old TeX, retaining the explicit equation ID/numbering flag."""
        from hrs_bridge.math_adapter import MathAdapter
        d=Document(self.docx)
        old=d._element.xpath('.//m:oMathPara')[-1]
        old.getparent().replace(old,MathAdapter(self.rt).omml('x+1',True))
        path=self.folder/'equation-edit.docx';d.save(path)
        out=self.folder/'equation-back'
        self.convert(self.rt,path,out,provenance=Path(self.receipt['provenance']))
        ir=self.rt.analyze((out/'report.md').read_bytes())
        eq=next(b for b in ir['blocks'] if b.get('id')=='eq-c7babca5e6cd')
        self.assertTrue(eq['numbered'])
        self.assertEqual(''.join(eq['latex'].split()),'x+1')

    def test_group_child_caption_edit_does_not_create_a_data_table(self):
        """B4: read the changed child caption, retain six resources and one group identity."""
        d=Document(self.docx)
        d.tables[-1].cell(1,0).paragraphs[0].text='(a) 从Word更新'
        path=self.folder/'group-edit.docx';d.save(path)
        out=self.folder/'group-back'
        self.convert(self.rt,path,out,provenance=Path(self.receipt['provenance']))
        ir=self.rt.analyze((out/'report.md').read_bytes())
        group=next(b for b in ir['blocks'] if b.get('id')=='fig-4c59dfec6c89')
        self.assertEqual(group['type'],'image-group')
        self.assertEqual(group['data']['images'][0]['caption'],'从Word更新')
        self.assertEqual(len(group['data']['images']),6)
        self.assertEqual(sum(b['type']=='table' for b in ir['blocks']),4)

    def test_legacy_combined_cells_still_read_edited_child_captions(self):
        """GROUP-02: retained exports with same-cell captions remain editable."""
        doc=Document(self.docx);table=doc.tables[-1]
        for row in range(len(table.rows)-2,-1,-2):
            pictures,captions=table.rows[row:row+2]
            for picture,caption in zip(pictures.cells,captions.cells):
                for paragraph in list(caption._tc.findall('{'+W+'}p')):picture._tc.append(paragraph)
            table._tbl.remove(captions._tr)
        table.cell(0,0).paragraphs[-1].text='(a) 旧布局仍可更新'
        path=self.folder/'legacy-group.docx';doc.save(path);out=self.folder/'legacy-group-back'
        self.convert(self.rt,path,out,provenance=Path(self.receipt['provenance']))
        ir=self.rt.analyze((out/'report.md').read_bytes(),out/'assets')
        group=next(b for b in ir['blocks'] if b.get('id')=='fig-4c59dfec6c89')
        self.assertEqual(group['data']['images'][0]['caption'],'旧布局仍可更新')
        self.assertEqual(len(group['data']['images']),6)

    def test_separate_caption_layout_refuses_unmapped_picture_or_image_text(self):
        """GROUP-02: splitting rows must not permit silently discarded content."""
        for kind in ('caption-picture','image-text'):
            with self.subTest(kind=kind):
                doc=Document(self.docx);table=doc.tables[-1]
                if kind=='caption-picture':
                    table.cell(1,0).paragraphs[0].add_run().add_picture(str(next((self.source/'assets').glob('*.png'))))
                    reason='caption row must not contain drawings|IMAGE_LAYOUT_EDIT_UNSUPPORTED'
                else:
                    table.cell(0,0).paragraphs[0].add_run('This text must not disappear')
                    reason='Unexpected text in group image row'
                path=self.folder/(kind+'.docx');doc.save(path);out=self.folder/(kind+'-back')
                with self.assertRaisesRegex(BridgeError,reason):
                    self.convert(self.rt,path,out,provenance=Path(self.receipt['provenance']))
                self.assertFalse(out.exists())

    def test_untracked_new_content_is_not_silently_dropped(self):
        """B4: insertion outside all mapped ranges requires a new mapping, never old-source restore."""
        d=Document(self.docx);d.add_paragraph('NEW_UNMAPPED_CONTENT')
        path=self.folder/'extra.docx';d.save(path)
        with self.assertRaisesRegex(BridgeError,'UNTRACKED_CONTENT'):
            self.convert(self.rt,path,self.folder/'no-extra',provenance=Path(self.receipt['provenance']))

    def test_external_unmarked_image_table_uses_confirmed_default_rule(self):
        """R17: image-bearing table without a table caption projects as a figure."""
        d=Document(ROOT/'fixtures/target_template.docx')
        table=d.add_table(rows=1,cols=2)
        picture=next((self.source/'assets').glob('*.png'))
        from docx.shared import Inches
        table.cell(0,0).paragraphs[0].add_run().add_picture(str(picture),width=Inches(1))
        path=self.folder/'unknown-grid.docx';d.save(path)
        self.convert(self.rt,path,self.folder/'image-grid')
        ir=self.rt.analyze((self.folder/'image-grid/report.md').read_bytes(),self.folder/'image-grid/assets')
        self.assertEqual(len([b for b in ir['blocks'] if b['type']=='figure']),1)

    def test_replaced_native_png_uses_current_bytes(self):
        """B4: replacement image bytes win over archived preview/original resources."""
        from io import BytesIO
        from PIL import Image
        from docx.oxml.ns import qn
        d=Document(self.docx);blob=BytesIO();Image.new('RGB',(48,32),'red').save(blob,format='PNG')
        rid=d.inline_shapes[0]._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
        d.part.related_parts[rid]._blob=blob.getvalue()
        path=self.folder/'image-edit.docx';d.save(path)
        out=self.folder/'image-back'
        self.convert(self.rt,path,out,provenance=Path(self.receipt['provenance']))
        ir=self.rt.analyze((out/'report.md').read_bytes(),out/'assets')
        figure=next(b for b in ir['blocks'] if b.get('id')=='fig-c6394cd828f7')
        self.assertEqual((out/figure['data']['src']).read_bytes(),blob.getvalue())

    def test_cropped_image_is_reported_instead_of_restoring_uncropped_source(self):
        """B4: a native edit outside the HRS schema must stop with a diagnostic."""
        from docx.oxml import OxmlElement
        d=Document(self.docx);fill=d.inline_shapes[0]._inline.xpath('.//pic:blipFill')[0]
        crop=OxmlElement('a:srcRect');crop.set('l','20000');fill.insert(1,crop)
        path=self.folder/'crop.docx';d.save(path)
        with self.assertRaisesRegex(BridgeError,'CROPPED_IMAGE_UNSUPPORTED'):
            self.convert(self.rt,path,self.folder/'no-crop',provenance=Path(self.receipt['provenance']))

    def test_unrepresented_native_image_size_edit_is_reported(self):
        """B4: a native resize cannot silently reuse a stale Word layout hint or old source."""
        d=Document(self.docx);d.inline_shapes[0].width+=91440
        path=self.folder/'resize.docx';d.save(path)
        with self.assertRaisesRegex(BridgeError,'IMAGE_LAYOUT_EDIT_UNSUPPORTED'):
            self.convert(self.rt,path,self.folder/'no-resize',provenance=Path(self.receipt['provenance']))

    def test_external_wps_sample_uses_real_heading_styles_and_math(self):
        """B4: Title and H1–H6, one data table/image and two OMML objects remain distinguishable."""
        out = self.folder / 'external'
        result = self.convert(self.rt, ROOT/'fixtures/wps_sample.docx', out)
        ir = self.rt.analyze((out/'report.md').read_bytes(), out/'assets')
        self.assertTrue(ir['ok'])
        levels = [b['level'] for b in ir['blocks'] if b['type']=='heading']
        self.assertEqual(levels, [1, 2, 3, 4, 5, 6])
        self.assertIn('- ',(out/'report.md').read_text('utf8'))  # user-approved H6 overflow projection
        self.assertEqual(sum(b['type']=='figure' for b in ir['blocks']), 1)
        self.assertEqual(sum(b['type']=='table' for b in ir['blocks']), 1)
        output = self.folder/'external-back.docx'
        to_docx(self.rt, out, output, accept_numbered_headings=True)
        doc = Document(output)
        self.assertEqual(len(doc._element.xpath('.//m:oMath')), 2)
        self.assertEqual(doc.paragraphs[0].style.name, 'Title')
        original=Document(ROOT/'fixtures/wps_sample.docx')
        self.assertEqual(doc.inline_shapes[0].width,original.inline_shapes[0].width)
        self.assertEqual(doc.inline_shapes[0].height,original.inline_shapes[0].height)
        self.assertTrue(result['diagnostics'])  # static source numbers are explicitly normalized


if __name__ == '__main__': unittest.main()
