"""FORMS-01: real CLI stages, article batches and refusal boundaries.

Fixtures are authored here. No inference responses are taken from a truth file.
"""
from pathlib import Path
from zipfile import ZipFile
import copy
import hashlib
import json
import os
import tempfile
import unittest
from lxml import etree as E
from docx import Document
from hrs_bridge import batch,forms
from hrs_bridge.runtime import BridgeError
from test_html import runtime

S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'


def fill(path,values):
    with ZipFile(path) as archive:parts={n:archive.read(n) for n in archive.namelist()}
    sheet=E.fromstring(parts['xl/worksheets/sheet1.xml'])
    for ref,value in values.items():
        cell=sheet.find('.//{'+S+'}c[@r="'+ref+'"]')
        if cell is None:raise AssertionError(ref)
        for child in list(cell):cell.remove(child)
        cell.set('t','inlineStr');E.SubElement(E.SubElement(cell,'{'+S+'}is'),'{'+S+'}t').text=str(value)
    parts['xl/worksheets/sheet1.xml']=E.tostring(sheet,xml_declaration=True,encoding='UTF-8')
    with ZipFile(path,'w') as archive:
        for name,raw in parts.items():archive.writestr(name,raw)


class FormsWorkflow(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix='bridge-forms-');self.home=Path(self.tmp.name);self.rt=runtime()
        self.inputs=self.home/'来源';self.inputs.mkdir();self.before={}
        for index in range(2):
            path=self.inputs/f'{index+1:02d}-报告.docx';doc=Document();doc.add_paragraph(f'试验报告{index+1}',style='Title')
            doc.add_paragraph('目的',style='Heading 1');doc.add_paragraph(f'独立测量{index+1}：2.30 V 与 1.5 mm 必须保持。')
            doc.add_paragraph('结果',style='Heading 1');doc.add_paragraph('保留这个完整结论。');doc.save(path)
            self.before[path]=path.read_bytes()
        self.run=self.home/'output'
    def tearDown(self):self.tmp.cleanup()
    def manifest(self,**extra):
        task={'id':'forms','source':{'kind':'wps-forms','path':str(self.inputs)},'outputs':['docx'],
              'forms_options':{'pages':'off','doc_conversion':'off'},**extra}
        path=self.home/'job.json';path.write_text(json.dumps({'schema':batch.SCHEMA,'output':str(self.run),'tasks':[task]}),encoding='utf8');return path
    def start(self):
        result=batch.run(self.rt,self.manifest());self.assertNotEqual(result['status'],'FAIL',result)
        self.assertEqual(result['integrity'],'PASS');task=result['tasks'][0]
        self.assertEqual(task['route'],'wps-forms');self.assertEqual(len(task['deliverables']),3)
        self.assertTrue(task['native_status']['required_steps_attempted']);return result
    def current(self):
        task=batch.status(self.rt,self.run)['tasks'][0];return self.run/task['forms_current'],task
    def test_first_draft_classification_two_feedback_batches_preserve_previous_outputs(self):
        self.start();first,task=self.current();old={self.run/d['path']:(self.run/d['path']).read_bytes() for d in task['deliverables']}
        form=first/'02_扫描核查.xlsx';fill(form,{'B6':'确认本表','F9':'归类A','H9':'展示甲'})
        self.assertEqual(batch.status(self.rt,self.run)['integrity'],'PASS')
        classified=forms.submit(self.rt,self.run,'forms','classify',form,'user');self.assertEqual(classified['integrity'],'PASS',classified)
        for number in (1,2):
            current,task=self.current();meta=batch.read(current/'feedback_export.json');row=9+len(meta['exports'])
            form=current/'03_反馈提交.xlsx'
            fill(form,{'B6':'提交本表',f'B{row}':'master',f'C{row}':f'独立测量{number}：2.30 V 与 1.5 mm 必须保持。',
                       f'D{row}':'字号',f'E{row}':str(13+number),f'G{row}':'仅此处',f'H{row}':'提交修改'})
            result=forms.submit(self.rt,self.run,'forms','feedback',form,'model');self.assertEqual(result['integrity'],'PASS',result)
        current,task=self.current();self.assertEqual(len(task['forms_rounds']),4)
        self.assertEqual([r['origin'] for r in task['forms_rounds']],['program_draft','user','model','model'])
        master=self.run/next(d['path'] for d in task['deliverables'] if d['document']=='master')
        document=Document(master)
        for number in (1,2):
            paragraph=next(p for p in document.paragraphs if p.text==f'独立测量{number}：2.30 V 与 1.5 mm 必须保持。')
            self.assertEqual(paragraph.runs[0].font.size.pt,13+number)
        for path,raw in {**self.before,**old}.items():self.assertEqual(path.read_bytes(),raw)
        self.assertTrue((self.run/'DELIVERY.md').is_file());self.assertTrue((self.run/'finish_list.json').is_file())
    def test_only_current_forms_are_mutable_and_old_workbook_is_refused(self):
        self.start();first,task=self.current();form=first/'02_扫描核查.xlsx';fill(form,{'B6':'确认本表'})
        forms.submit(self.rt,self.run,'forms','classify',form)
        with self.assertRaises(BridgeError):forms.submit(self.rt,self.run,'forms','classify',form)
        self.assertEqual(batch.status(self.rt,self.run)['integrity'],'PASS')
        current,task=self.current();out=self.run/task['deliverables'][0]['path'];out.write_bytes(out.read_bytes()+b'tampered')
        self.assertEqual(batch.status(self.rt,self.run)['integrity'],'FAIL')
        with self.assertRaises(BridgeError):forms.submit(self.rt,self.run,'forms','feedback',current/'03_反馈提交.xlsx')
    def test_changed_runtime_original_or_frozen_input_is_refused(self):
        import shutil
        self.start();current,task=self.current()
        other=self.home/'other-core';shutil.copytree(self.rt.forms,other)
        target=other/'forms_tool.py';target.write_bytes(target.read_bytes()+b'\n# changed runtime\n')
        changed=copy.copy(self.rt);changed.forms=other
        with self.assertRaisesRegex(BridgeError,'code changed'):batch.status(changed,self.run)
        source=next(iter(self.before));source.write_bytes(source.read_bytes()+b'changed')
        with self.assertRaisesRegex(BridgeError,'Source changed'):forms.submit(self.rt,self.run,'forms','feedback',current/'03_反馈提交.xlsx')
        source.write_bytes(self.before[source]);frozen=Path(task['inputs'][-1]['frozen']);frozen.write_bytes(frozen.read_bytes()+b'changed')
        self.assertEqual(batch.status(self.rt,self.run)['integrity'],'FAIL')
    def test_article_ordered_review_plan_is_bound_and_does_not_split_an_article(self):
        for path in self.before:
            doc=Document();doc.add_paragraph(path.stem,style='Title');doc.add_paragraph('1 FFT频谱分析',style='Heading 1')
            doc.add_paragraph('这里是完整技术内容，保持 2.30 V。');doc.save(path)
        self.start();current,task=self.current();before=batch.inventory(self.run)
        plan=forms.review_plan(self.rt,self.run,'forms',.2)
        self.assertGreater(plan['pending'],0)
        self.assertEqual(batch.inventory(self.run),before);self.assertIn('feedback_export_sha256',plan['bindings'])
        ids=[a['source_id'] for group in plan['batches'] for a in group['articles']]
        self.assertEqual(ids,sorted(ids));self.assertEqual(len(ids),len(set(ids)))
        self.assertEqual(ids,['S0001','S0002'])
        self.assertEqual(sum(group['pending'] for group in plan['batches']),plan['pending'])
        for fraction in (0,1.1,float('nan')):
            with self.assertRaises(BridgeError):forms.review_plan(self.rt,self.run,'forms',fraction)
    def test_route_options_are_strict(self):
        for extra in ({'review_policy':'pause'},{'forms_options':{'policy':'接受程序建议'}},{'word_options':{}},{'forms_options':{'invent':True}}):
            value=json.loads(self.manifest(**extra).read_text())
            with self.assertRaises(BridgeError):batch.validate_manifest(value)
    def test_workbook_changed_during_freeze_is_rejected_and_next_submit_can_continue(self):
        from unittest.mock import patch
        self.start();current,task=self.current();form=current/'02_扫描核查.xlsx';fill(form,{'B6':'确认本表'})
        original=Path.read_bytes;calls=0
        def changing(path):
            nonlocal calls
            raw=original(path)
            if path==form:
                calls+=1
                if calls==2:return raw+b'changed concurrently'
            return raw
        with patch.object(Path,'read_bytes',changing):
            with self.assertRaisesRegex(BridgeError,'while freezing'):
                forms.submit(self.rt,self.run,'forms','classify',form)
        result=batch.status(self.rt,self.run);self.assertEqual(result['integrity'],'PASS')
        self.assertEqual(result['tasks'][0]['forms_attempts'][-1]['status'],'REJECTED')
        self.assertEqual(len(result['tasks'][0]['forms_rounds']),1)
        result=forms.submit(self.rt,self.run,'forms','classify',form)
        self.assertEqual(result['integrity'],'PASS');self.assertEqual(len(result['tasks'][0]['forms_rounds']),2)
    def test_editable_declaration_cannot_unlock_a_document(self):
        self.start();current,task=self.current();record=batch.read(self.run/'WORKFLOW.json')
        output=task['deliverables'][0]['path'];(self.run/output).write_bytes((self.run/output).read_bytes()+b'bad')
        record['tasks'][0]['editable_forms'].append(str(Path(output).relative_to('tasks/forms')).replace('\\','/'))
        batch.write(self.run/'WORKFLOW.json',record)
        self.assertEqual(batch.status(self.rt,self.run)['integrity'],'FAIL')
    def assert_failed_update_can_be_resubmitted(self,operation,fault):
        from unittest.mock import patch
        self.start();current,task=self.current()
        form=current/('02_扫描核查.xlsx' if operation=='classify' else '03_反馈提交.xlsx')
        if operation=='classify':fill(form,{'B6':'确认本表'})
        else:
            row=9+len(batch.read(current/'feedback_export.json')['exports'])
            fill(form,{'B6':'提交本表',f'B{row}':'master',f'C{row}':'独立测量1：2.30 V 与 1.5 mm 必须保持。',
                       f'D{row}':'字号',f'E{row}':'14',f'G{row}':'仅此处',f'H{row}':'提交修改'})
        keys=('route','status','forms_current','editable_forms','deliverables','native_status','next_action',
              'pending_count','source_gaps','forms_rounds','source_order')
        committed={key:copy.deepcopy(task[key]) for key in keys}
        old={self.run/d['path']:(self.run/d['path']).read_bytes() for d in task['deliverables']}
        hits=[];round_name='002-'+operation
        if fault=='native-lock':
            import ctypes
            from ctypes import wintypes
            kernel=ctypes.WinDLL('kernel32',use_last_error=True)
            kernel.CreateFileW.argtypes=[wintypes.LPCWSTR,wintypes.DWORD,wintypes.DWORD,wintypes.LPVOID,wintypes.DWORD,wintypes.DWORD,wintypes.HANDLE]
            kernel.CreateFileW.restype=wintypes.HANDLE;kernel.CloseHandle.argtypes=[wintypes.HANDLE]
            original=batch.fingerprint
            def locked(path):
                if not hits and Path(path).name=='master.docx' and round_name in Path(path).parts:
                    handle=kernel.CreateFileW(str(path),0x80000000,0,None,3,0x80,None)
                    if handle==ctypes.c_void_p(-1).value:raise ctypes.WinError(ctypes.get_last_error())
                    hits.append(str(path))
                    try:return original(path)
                    finally:kernel.CloseHandle(handle)
                return original(path)
            injection=patch.object(batch,'fingerprint',locked);expected=PermissionError
        elif fault=='metadata':
            original=batch.read
            def unavailable(path):
                if not hits and Path(path).name=='feedback_export.json' and round_name in Path(path).parts:
                    hits.append(str(path));raise PermissionError('One transient metadata read failure')
                return original(path)
            injection=patch.object(batch,'read',unavailable);expected=PermissionError
        else:
            original=forms.command
            def corrupt(rt,home,op,*args):
                value=original(rt,home,op,*args)
                if op==operation and not hits:
                    dest=Path(args[args.index('--out')+1]);state=batch.read(dest/'forms_state.json')
                    core=dest/state['drafts'][-1]['run'];item=next(d for d in value['latest']['deliverables'] if d['id']=='master')
                    path=core/item['file'];path.write_bytes(path.read_bytes()+b'wrong digest');hits.append(str(path))
                return value
            injection=patch.object(forms,'command',corrupt);expected=BridgeError
        with injection:
            with self.assertRaises(expected):forms.submit(self.rt,self.run,'forms',operation,form)
        self.assertEqual(len(hits),1)
        failed=batch.status(self.rt,self.run);self.assertEqual(failed['integrity'],'PASS',failed)
        task=failed['tasks'][0];self.assertEqual({key:task[key] for key in keys},committed)
        self.assertEqual(task['forms_attempts'][-1]['status'],'REJECTED')
        attempted=self.run/task['forms_attempts'][-1]['output'];self.assertTrue(list(attempted.rglob('master.docx')))
        result=forms.submit(self.rt,self.run,'forms',operation,form)
        self.assertEqual(result['integrity'],'PASS',result);self.assertEqual(result['status'],'COMPLETE',result)
        task=result['tasks'][0];self.assertTrue(task['forms_current'].endswith('003-'+operation))
        self.assertEqual(len(task['forms_rounds']),len(committed['forms_rounds'])+1)
        self.assertEqual([a['status'] for a in task['forms_attempts']],['REJECTED','APPLIED'])
        self.assertTrue(attempted.is_dir())
        for path,raw in {**self.before,**old}.items():self.assertEqual(path.read_bytes(),raw)
    @unittest.skipUnless(os.name=='nt','Windows file sharing semantics')
    def test_classify_output_read_lock_preserves_committed_state_and_allows_resubmit(self):
        self.assert_failed_update_can_be_resubmitted('classify','native-lock')
    @unittest.skipUnless(os.name=='nt','Windows file sharing semantics')
    def test_feedback_output_read_lock_preserves_committed_state_and_allows_resubmit(self):
        self.assert_failed_update_can_be_resubmitted('feedback','native-lock')
    def test_late_metadata_read_failure_preserves_committed_state_and_allows_resubmit(self):
        self.assert_failed_update_can_be_resubmitted('classify','metadata')
    def test_invalid_new_deliverable_is_rejected_without_committing_current_round(self):
        self.assert_failed_update_can_be_resubmitted('classify','bad-sha')
    def test_source_gap_is_local_and_independent_legacy_task_still_finishes(self):
        (self.inputs/'bad.docx').write_bytes(b'not a docx')
        value=json.loads(self.manifest().read_text());value['tasks'].append(dict(id='legacy',source=dict(kind='docx',path=str(next(iter(self.before)))),outputs=['docx']))
        job=self.home/'job.json';job.write_text(json.dumps(value),encoding='utf8');result=batch.run(self.rt,job)
        self.assertEqual(result['integrity'],'PASS',result)
        self.assertEqual(result['tasks'][0]['status'],'PARTIAL',result)
        self.assertTrue(result['tasks'][0]['deliverables'])
        self.assertTrue(result['tasks'][0]['source_gaps'])
        self.assertEqual(result['tasks'][1]['status'],'COMPLETE',result)


if __name__=='__main__':unittest.main()
