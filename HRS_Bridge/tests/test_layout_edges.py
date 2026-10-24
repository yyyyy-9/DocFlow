"""Small missing geometry cases, using current HRS serializers and WPS formatting."""
import json
from pathlib import Path
import tempfile
import unittest
from PIL import Image
from docx import Document
from docx.oxml.ns import qn
from hrs_bridge.browser import call
from hrs_bridge.docx_export import convert as export
from hrs_bridge.docx_import import convert as restore
from hrs_bridge.ooxml import N, W, zero_indent
from hrs_bridge.runtime import BridgeError
from test_html import runtime


class LayoutEdges(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory(prefix='bridge-layout-test-')
        cls.root=Path(cls.tmp.name);cls.rt=runtime();cls.source=cls.root/'source'
        (cls.source/'assets').mkdir(parents=True)
        Image.new('RGB',(1200,60),'white').save(cls.source/'assets/wide.png')
        Image.new('RGB',(60,1200),'white').save(cls.source/'assets/tall.png')
        cls.groups=[]
        for ident,rows,cols,count,labels in [('fig-wide',1,4,4,False),('fig-tall',3,4,12,True),('fig-ragged',2,3,5,True)]:
            cls.groups.append({'schema':'report-image-group-v1','id':ident,'rows':rows,'columns':cols,
                'caption':'版式边界','showSubfigureLabels':labels,
                'images':[{'src':'assets/'+('wide.png' if i%2 else 'tall.png'),'caption':'长说明，保留顺序与文字。'*5} for i in range(count)]})
        snippets=call(cls.rt,'serialize',items=[{'type':'image-group','data':g} for g in cls.groups])
        cell=lambda text,rs=1,cs=1:dict(text=text,rowspan=rs,colspan=cs)
        cls.table={'schema':'report-table-v1','id':'tbl-merged','caption':'合并单元格',
            'rows':[[cell('A',2),cell('B',1,2),None],[None,cell('C'),cell('D')]]}
        (cls.source/'report.md').write_text('# 边界\n\n正文仍使用模板缩进。A &amp; B\n\n'+'\n'.join(snippets)+
            '\n| A | B |\n| --- | --- |\n| 1 | 2 |\n\n```text\nliteral &amp;\n```\n\n'+
            '```report-table\n'+json.dumps(cls.table,ensure_ascii=False)+'\n```\n',encoding='utf8')
        cls.output=cls.root/'edge.docx';cls.receipt=export(cls.rt,cls.source,cls.output)

    @classmethod
    def tearDownClass(cls):cls.tmp.cleanup()

    def test_ragged_grid_labels_and_bounded_images(self):
        """1×4, 3×4 and partial last row retain capacity and label policy, without fixed heights."""
        doc=Document(self.output)
        self.assertIn('正文仍使用模板缩进。A & B',[p.text for p in doc.paragraphs])
        self.assertIn('literal &amp;\n',[p.text for p in doc.paragraphs])
        for table,data in zip(doc.tables,self.groups):
            self.assertEqual((len(table.rows),len(table.columns)),(data['rows']*2,data['columns']))
            self.assertEqual(len(table._tbl.xpath('.//w:drawing')),len(data['images']))
            self.assertFalse(table._tbl.xpath('.//w:trHeight|.//w:tblHeader'))
            captions=[p.text for row in table.rows for c in row.cells for p in c.paragraphs if p.style.name=='图题']
            self.assertEqual(captions[0].startswith('(a)'),data['showSubfigureLabels'])
            for row in table.rows[1::2]:
                self.assertFalse(row._tr.xpath('.//w:drawing'))
                self.assertTrue(all(c._tc.xpath('./w:tcPr/w:vAlign/@w:val')==['top'] for c in row.cells))
            for p in table._tbl.xpath('.//w:p'):
                self.assertEqual(p.find('w:pPr/w:ind',N).get(qn('w:firstLineChars')),'0')
        out=self.root/'exact'
        restore(self.rt,self.output,out)
        self.assertEqual((out/'report.md').read_bytes(),(self.source/'report.md').read_bytes())

    def test_merged_table_edit_retains_native_span_structure(self):
        """An edited anchor keeps both the horizontal and vertical merge contracts."""
        doc=Document(self.output);doc.tables[-1].cell(0,0).text='NEW'
        edited=self.root/'merged-edit.docx';doc.save(edited)
        out=self.root/'merged-back'
        restore(self.rt,edited,out,provenance=Path(self.receipt['provenance']))
        ir=self.rt.analyze((out/'report.md').read_bytes(),out/'assets')
        table=next(b for b in ir['blocks'] if b.get('id')=='tbl-merged')['data']
        self.assertEqual(table['rows'][0][0],dict(text='NEW',rowspan=2,colspan=1))
        self.assertEqual(table['rows'][0][1]['colspan'],2)
        self.assertIsNone(table['rows'][0][2]);self.assertIsNone(table['rows'][1][0])

    def test_partial_last_row_reads_caption_at_its_logical_image_position(self):
        """GROUP-02: the second caption row maps to child e, without a sixth child."""
        doc=Document(self.output);doc.tables[2].cell(3,1).paragraphs[0].text='(e) Updated fifth image'
        path=self.root/'fifth-edit.docx';doc.save(path);out=self.root/'fifth-back'
        restore(self.rt,path,out,provenance=Path(self.receipt['provenance']))
        ir=self.rt.analyze((out/'report.md').read_bytes(),out/'assets')
        group=next(b['data'] for b in ir['blocks'] if b.get('id')=='fig-ragged')
        self.assertEqual(len(group['images']),5)
        self.assertEqual(group['images'][4]['caption'],'Updated fifth image')

    def test_empty_last_caption_cell_rejects_new_unmapped_text(self):
        """GROUP-02: an unused grid slot cannot absorb untracked new text."""
        doc=Document(self.output);doc.tables[2].cell(3,2).text='Keep this new content'
        path=self.root/'empty-edit.docx';doc.save(path);out=self.root/'empty-back'
        with self.assertRaisesRegex(BridgeError,'Unexpected group content in empty cell'):
            restore(self.rt,path,out,provenance=Path(self.receipt['provenance']))
        self.assertFalse(out.exists())


class IndentGuard(unittest.TestCase):
    def test_wps_extent_serialization_is_not_a_native_resize(self):
        """B5: WPS truncates fractional twips; a real point-size edit must still differ."""
        from docx.oxml import OxmlElement
        from hrs_bridge.ooxml import drawing_extents
        before=OxmlElement('wp:inline');extent=OxmlElement('wp:extent')
        extent.set('cx',str(8312*635));extent.set('cy',str(3880*635+600));before.append(extent)
        after=OxmlElement('wp:inline');saved=OxmlElement('wp:extent')
        saved.set('cx',str(8312*635));saved.set('cy',str(3880*635));after.append(saved)
        self.assertEqual(drawing_extents([before]),drawing_extents([after]))
        saved.set('cy',str(3900*635))
        self.assertNotEqual(drawing_extents([before]),drawing_extents([after]))

    def test_first_line_chars_is_not_left_inherited(self):
        """Clearing twips alone must not leave the template's character indentation in force."""
        doc=Document();p=doc.add_paragraph()
        from hrs_bridge.ooxml import element
        p._p.get_or_add_pPr().append(element('w:ind',firstLine=0,firstLineChars=200,leftChars=100))
        zero_indent(p._p)
        self.assertEqual(set(p._p.find('w:pPr/w:ind',N).attrib.values()),{'0'})


if __name__=='__main__':unittest.main()
