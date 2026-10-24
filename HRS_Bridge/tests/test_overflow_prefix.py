"""FB-04: the overflow list contains title content, not a second chapter label."""
import json
from pathlib import Path
import tempfile
import unittest
from docx import Document
from hrs_bridge.docx_import import convert
from hrs_bridge.docx_export import convert as to_word
from test_html import runtime


class OverflowPrefix(unittest.TestCase):
    def test_six_level_literal_prefix_is_removed_and_word_role_survives(self):
        """A confirmed Heading 6 becomes a clean list; bold text and its native role survive."""
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp);rt=runtime();doc=Document();doc.add_paragraph('报告题名',style='Title')
            for level in range(1,6):doc.add_paragraph('父级'+str(level),style='Heading '+str(level))
            p=doc.add_paragraph(style='Heading 6');p.add_run('1.1.1.1.1.1 五级标题示例').bold=True
            source=folder/'source.docx';doc.save(source);convert(rt,source,folder/'back')
            text=(folder/'back/report.md').read_text('utf8')
            self.assertIn('- **五级标题示例**',text)
            self.assertNotIn('1.1.1.1.1.1',text)
            origin=json.loads((folder/'back/bridge-origin.json').read_text('utf8'))
            self.assertEqual(origin['word_styles'][-1]['source_number_prefix'],'1.1.1.1.1.1 ')
            to_word(rt,folder/'back',folder/'again.docx')
            p=Document(folder/'again.docx').paragraphs[-1]
            self.assertEqual((p.style.name,p.text),('Heading 6','五级标题示例'))
            self.assertTrue(p.runs[0].bold)

    def test_quantities_and_unnumbered_titles_are_preserved(self):
        """A leading measurement/year is content; no chapter-prefix match means no deletion."""
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp);rt=runtime();doc=Document()
            for text in ('1.5 mm 厚度范围','2026 年总结','无编号小标题'):doc.add_paragraph(text,style='Heading 6')
            source=folder/'source.docx';doc.save(source);convert(rt,source,folder/'back')
            text=(folder/'back/report.md').read_text('utf8')
            for expected in ('- 1.5 mm 厚度范围','- 2026 年总结','- 无编号小标题'):self.assertIn(expected,text)

    def test_number_like_body_text_is_not_rewritten(self):
        """The rule applies only to a confirmed heading being projected as a list."""
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp);rt=runtime();doc=Document();doc.add_paragraph('正文中的版本路径为 1.1.1.1.1.1，必须保留。')
            source=folder/'source.docx';doc.save(source);convert(rt,source,folder/'back')
            self.assertIn('版本路径为 1.1.1.1.1.1',(folder/'back/report.md').read_text('utf8'))


if __name__=='__main__':unittest.main()
