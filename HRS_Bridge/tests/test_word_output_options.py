"""Native footnotes/references are derived Word output; Markdown stays exact."""
from pathlib import Path
import json,shutil,tempfile,unittest,zipfile
from docx import Document
from lxml import etree as E
from hrs_bridge.docx_export import convert,prepare
from hrs_bridge.docx_import import convert as restore
from hrs_bridge.browser import call
from hrs_bridge.runtime import BridgeError
from hrs_bridge.ooxml import N
from test_html import runtime

class WordOptions(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory(prefix='word-policy-');cls.root=Path(cls.tmp.name);cls.rt=runtime()
        cls.project=cls.root/'source';(cls.project/'assets').mkdir(parents=True)
        cls.raw='# 输出策略\n\n> 摘要：这里是正常的摘要正文。\n\n## 第一章\n\n正文**强调**与*中文斜体*，公式 $\\boldsymbol{E}=mc^2$，来源[^a]，重复[^a]。\n\n> 第一行\n> 第二行\n\n## 第二章\n\n另一个来源[^b]。\n\n[^a]: 文献**一**，公式 $x^2$，以及[链接](https://example.com/)。\n\n[^b]: 文献二。\n'
        (cls.project/'report.md').write_text(cls.raw,'utf8',newline='')

    @classmethod
    def tearDownClass(cls):cls.tmp.cleanup()

    def test_01_reference_section_cover_toc_and_exact_return(self):
        template=self.rt.wps/'templates/cover_template_v0311.docx';before=template.read_bytes()
        target=self.root/'references.docx';convert(self.rt,self.project,target,template=template,options={'notes_mode':'references','preserve_emphasis':False})
        doc=Document(target);heading=next(p for p in doc.paragraphs if p.text=='参考文献')
        self.assertEqual(heading._p.xpath('./w:pPr/w:numPr/w:numId/@w:val'),['0'])
        self.assertEqual(heading._p.xpath('./w:pPr/w:outlineLvl/@w:val'),['0'])
        self.assertTrue(doc._element.xpath('//w:instrText[contains(.,"TOC")]|//w:fldSimple[contains(@w:instr,"TOC")]'))
        self.assertEqual(len(doc._element.xpath('//w:fldSimple[contains(@w:instr,"SEQ HRSReference")]')),2)
        self.assertEqual(len(doc._element.xpath('//w:fldSimple[contains(@w:instr," REF HR_")]')),3)
        self.assertFalse(doc._element.xpath('//w:footnoteReference'))
        with zipfile.ZipFile(template) as old,zipfile.ZipFile(target) as new:
            for name in old.namelist():
                if name.startswith(('word/header','word/footer')):self.assertEqual(old.read(name),new.read(name))
        restored=self.root/'returned';restore(self.rt,target,restored)
        self.assertEqual((restored/'report.md').read_text('utf8'),self.raw);self.assertEqual(template.read_bytes(),before)
        changed=self.root/'changed-references.docx'
        with zipfile.ZipFile(target) as old,zipfile.ZipFile(changed,'x',zipfile.ZIP_DEFLATED) as new:
            for item in old.infolist():
                data=old.read(item.filename)
                if item.filename=='word/document.xml':data=data.replace('文献二。'.encode(),'被编辑的来源。'.encode())
                new.writestr(item,data)
        with self.assertRaisesRegex(BridgeError,'REFERENCE_EDIT_UNSUPPORTED'):
            restore(self.rt,changed,self.root/'rejected',provenance=Path(str(target)+'.hrs-bridge.json'))

    def test_02_abstract_label_emphasis_and_math_have_separate_contracts(self):
        math=[]
        for keep in (False,True):
            target=self.root/('keep.docx' if keep else 'strip.docx');convert(self.rt,self.project,target,options={'preserve_emphasis':keep})
            doc=Document(target);abstract=next(p for p in doc.paragraphs if p.text.startswith('摘要：'))
            self.assertTrue(abstract.runs[0].bold);self.assertTrue(all(r.bold is not True for r in abstract.runs[1:]))
            body=next(p for p in doc.paragraphs if p.text.startswith('正文'))
            self.assertEqual(any(r.bold for r in body.runs),keep);self.assertEqual(any(r.italic for r in body.runs),keep)
            math.append([E.tostring(n,method='c14n') for n in doc._element.xpath('//m:oMath')])
            quote=next(p for p in doc.paragraphs if p.text.startswith('第一行'));self.assertIn('\n第二行',quote.text)
            self.assertEqual(len(doc._element.xpath('//w:footnoteReference')),2)
            self.assertIn('[',body.text);self.assertIn(']',body.text)
        self.assertTrue(math[0]);self.assertEqual(math[0],math[1])

    def test_03_model_is_bound_to_output_options_and_invalid_options_fail(self):
        req=prepare(self.rt,self.project,options={'notes_mode':'references'})
        model=call(self.rt,'model',project=self.project,**req)
        with self.assertRaisesRegex(BridgeError,'快照不一致'):
            convert(self.rt,self.project,self.root/'wrong-policy.docx',model=model,options={'notes_mode':'footnotes'})
        self.assertFalse((self.root/'wrong-policy.docx').exists())
        for invalid in [{'notes_mode':'unknown'},{'preserve_emphasis':'false'},{'extra':True}]:
            with self.assertRaises(BridgeError):prepare(self.rt,self.project,options=invalid)

    def test_04_group_caption_math_remains_native_and_source_is_exact(self):
        from PIL import Image
        folder=self.root/'group';(folder/'assets').mkdir(parents=True)
        for name in ['a.png','b.png']:Image.new('RGB',(20,10),'blue').save(folder/'assets'/name)
        raw='# 组图公式\n\n## 内容\n\n::: {#fig-native layout-ncol=2}\n\n![子图 $x^2$](assets/a.png)\n\n![第二张](assets/b.png)\n\n组图 $E=mc^2$\n\n:::\n'
        (folder/'report.md').write_text(raw,'utf8',newline='');target=self.root/'group.docx';convert(self.rt,folder,target)
        doc=Document(target);self.assertEqual(len(doc._element.xpath('//m:oMath')),2);self.assertEqual(len(doc.tables),1)
        restored=self.root/'group-return';restore(self.rt,target,restored);self.assertEqual((restored/'report.md').read_text('utf8'),raw)

    def test_05_external_caption_text_does_not_invent_markdown_semantics(self):
        from PIL import Image
        from docx.shared import Inches
        from hrs_bridge.math_adapter import MathAdapter
        folder=self.root/'external-captions';folder.mkdir();image=folder/'sample.png'
        Image.new('RGB',(80,40),'navy').save(image)
        literal='$E=mc^2$ [[此处]] @fig:missing ==重点=='
        doc=Document();doc.add_paragraph('外来题注',style='Title')
        table=doc.add_table(rows=1,cols=2)
        for cell in table.rows[0].cells:
            cell.paragraphs[0].add_run().add_picture(str(image),width=Inches(1))
            cell.add_paragraph(literal)
        cap=doc.add_paragraph('图 1 '+literal+' 原生公式：')
        cap._p.append(MathAdapter(self.rt).omml('x^2',False))
        target=folder/'external.docx';doc.save(target);restore(self.rt,target,folder/'back')
        ir=self.rt.analyze((folder/'back/report.md').read_bytes(),folder/'back/assets')
        group=next(b['data'] for b in ir['blocks'] if b['type']=='image-group')
        self.assertIn('\\$E\\=mc^2\\$',group['caption']);self.assertIn('$x^{2}$',group['caption'])
        for item in group['images']:
            self.assertIn('\\$E\\=mc^2\\$',item['caption']);self.assertIn('\\@fig:missing',item['caption'])
            self.assertIn('\\[\\[此处\\]\\]',item['caption']);self.assertIn('\\=\\=重点\\=\\=',item['caption'])

if __name__=='__main__':unittest.main()
