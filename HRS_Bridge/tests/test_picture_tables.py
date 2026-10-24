"""External image tables use the shared WPS figure semantics, without provenance."""
import json
from pathlib import Path
import tempfile
import unittest
from docx import Document
from docx.shared import Inches
from PIL import Image
from hrs_bridge.docx_import import convert
from hrs_bridge.runtime import BridgeError
from test_html import runtime


class PictureTables(unittest.TestCase):
    def sample(self,folder,rows=2,columns=3,table_caption=False):
        image=folder/'image.png';Image.new('RGB',(120,60),'navy').save(image)
        doc=Document();doc.add_paragraph('图片表格',style='Title')
        if table_caption:doc.add_paragraph('表 3 原生表格')
        t=doc.add_table(rows=rows,cols=columns)
        for i in range(rows*columns):
            c=t.cell(i//columns,i%columns);c.paragraphs[0].add_run().add_picture(str(image),width=Inches(1))
            c.add_paragraph('工况 '+str(i+1)+'，量值 1.25')
        if not table_caption:doc.add_paragraph('图 9 六个工况')
        path=folder/'external.docx';doc.save(path);return path
    def test_external_two_by_three_table_becomes_one_group(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);rt=runtime();source=self.sample(p);original=source.read_bytes()
            receipt=convert(rt,source,p/'back')
            ir=rt.analyze((p/'back/report.md').read_bytes(),p/'back/assets')
            groups=[b for b in ir['blocks'] if b['type']=='image-group']
            self.assertEqual(len(groups),1);data=groups[0]['data']
            self.assertEqual((data['rows'],data['columns']),(2,3))
            self.assertEqual([im['caption'] for im in data['images']],['工况 '+str(i+1)+'，量值 1.25' for i in range(6)])
            self.assertEqual(data['caption'],'六个工况')
            self.assertFalse([b for b in ir['blocks'] if b['type']=='table'])
            self.assertEqual(receipt['recognition']['unresolved'],0)
            self.assertEqual(source.read_bytes(),original)
    def test_single_picture_table_is_single_figure(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);rt=runtime();convert(rt,self.sample(p,1,1),p/'back')
            ir=rt.analyze((p/'back/report.md').read_bytes(),p/'back/assets')
            self.assertEqual(len([b for b in ir['blocks'] if b['type']=='figure']),1)
            self.assertIn('工况 1，量值 1.25',(p/'back/report.md').read_text('utf8'))
    def test_caption_migrates_entire_old_ordinal_but_preserves_title_quantity(self):
        from hrs_bridge.docx_import import escaped
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);rt=runtime();source=self.sample(p,1,2);doc=Document(source)
            doc.paragraphs[-1].text='图2.1.3 1.25 mm 工况';doc.save(source);before=source.read_bytes()
            convert(rt,source,p/'back');ir=rt.analyze((p/'back/report.md').read_bytes(),p/'back/assets')
            group=next(b['data'] for b in ir['blocks'] if b['type']=='image-group')
            self.assertEqual(group['caption'],escaped('1.25 mm 工况'));self.assertEqual(source.read_bytes(),before)
    def test_separate_caption_row_stays_with_its_image_column(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);rt=runtime();source=self.sample(p,2,2);doc=Document(source)
            from docx.oxml.ns import qn
            for cell in doc.tables[0].rows[1].cells:
                for drawing in cell._tc.findall('.//'+qn('w:drawing')):drawing.getparent().remove(drawing)
            doc.save(source);convert(rt,source,p/'back')
            group=next(b['data'] for b in rt.analyze((p/'back/report.md').read_bytes(),p/'back/assets')['blocks'] if b['type']=='image-group')
            self.assertEqual((group['rows'],group['columns']),(1,2))
            self.assertEqual(group['images'][0]['caption'],'工况 1，量值 1.25 工况 3，量值 1.25')
    def test_unrepresentable_merged_grid_fails_without_partial_project(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);rt=runtime();source=self.sample(p,1,2);doc=Document(source)
            doc.tables[0].cell(0,0).merge(doc.tables[0].cell(0,1));doc.save(source)
            with self.assertRaisesRegex(BridgeError,'HRS_IMAGE_TABLE_LAYOUT_UNSUPPORTED'):convert(rt,source,p/'back')
            self.assertFalse((p/'back').exists())
    def test_native_word_reassembly_then_hrs_keeps_group_and_real_ref(self):
        from hrs_bridge.browser import call
        from hrs_bridge.docx_export import convert as to_docx
        from wps_pipeline.workflow import start_job,finish,status
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);rt=runtime();project=p/'project';(project/'assets').mkdir(parents=True)
            Image.new('RGB',(120,60),'navy').save(project/'assets/image.png')
            data=dict(schema='report-image-group-v1',id='fig-original',rows=2,columns=3,showSubfigureLabels=False,
                      caption='六图整体',images=[dict(src='assets/image.png',caption='子图 '+str(i+1)) for i in range(6)])
            source=call(rt,'serialize',items=[dict(type='image-group',data=data)])[0]
            (project/'report.md').write_text('# 六图报告\n\n## 对比\n\n'+source+'\n见 @fig:fig-original ，这里保留真实对象引用与完整的普通正文说明。\n','utf8')
            word=p/'source.docx';to_docx(rt,project,word)
            run=p/'native';state=start_job({'sources':[dict(id='src',path=str(word))],'make_master':False},p,run)
            self.assertFalse(state['review_requests']);finish(run,state)
            self.assertEqual(status(run,state)['status'],'COMPLETED_LOCAL_CHECKS',json.dumps(state.get('finish_list'),ensure_ascii=False))
            convert(rt,run/'build_001/documents/src.docx',p/'back')
            ir=rt.analyze((p/'back/report.md').read_bytes(),p/'back/assets')
            group=next(b for b in ir['blocks'] if b['type']=='image-group')
            self.assertEqual(len(group['data']['images']),6)
            self.assertEqual(group['data']['caption'],'六图整体')
            self.assertIn('@fig:'+group['id'],(p/'back/report.md').read_text('utf8'))


if __name__=='__main__':unittest.main()
