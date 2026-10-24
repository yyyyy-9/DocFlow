"""FLOW-01/02: real native route, isolated tasks and source-bound review."""
import json,os,sys,tempfile,unittest
from pathlib import Path
from docx import Document
from hrs_bridge.workflow import run
from test_html import runtime

class BatchWorkflow(unittest.TestCase):
 def setUp(self):self.tmp=tempfile.TemporaryDirectory();self.folder=Path(self.tmp.name);self.rt=runtime()
 def tearDown(self):self.tmp.cleanup()
 def document(self,name='source.docx',ambiguous=False):
  path=self.folder/name;doc=Document()
  doc.add_paragraph('1 FFT频谱分析' if ambiguous else '1 试验方法',style='Heading 1')
  doc.add_paragraph('这里保留完整内容，测量值 1.2毫米不得改成章节编号。');doc.save(path);return path
 def manifest(self,tasks,**extra):
  path=self.folder/'job.json';path.write_text(json.dumps({'schema':'hrs-bridge-job-v2','output':'output','tasks':tasks,**extra},ensure_ascii=False),encoding='utf8');return path
 def task(self,ident,source,**extra):return {'id':ident,'source':{'kind':'docx','path':str(source)},'outputs':['docx'],**extra}
 def test_native_word_word_finishes_all_existing_stages_without_markdown(self):
  src=self.document();raw=src.read_bytes();result=run(self.rt,self.manifest([self.task('word',src)]))
  self.assertEqual(result['status'],'COMPLETE',result);self.assertTrue(result['accepted'])
  task=result['tasks'][0];self.assertEqual(task['route'],'wps-native')
  self.assertEqual(task['native_status']['receipt_integrity'],'PASS')
  self.assertTrue(task['native_status']['required_steps_attempted']);self.assertEqual(src.read_bytes(),raw)
  out=self.folder/'output';self.assertTrue((out/'DELIVERY.md').is_file());self.assertTrue((out/'finish_list.json').is_file())
  self.assertFalse(list(out.rglob('report.md')))
  self.assertTrue(any(x['name']=='word.status' for x in task['stages']))
 def test_native_master_uses_shared_writer_after_batched_review(self):
  import importlib.util
  from hrs_bridge.batch import review,resume,status
  from wps_pipeline.verify_report_build import Package,checked_fields
  from wps_pipeline.verify_native_numbering import labels
  spec=importlib.util.spec_from_file_location('native_fixture',self.rt.wps/'tests/run_native_numbering.py')
  fixture=importlib.util.module_from_spec(spec);spec.loader.exec_module(fixture)
  src,blueprint=fixture.source(self.folder);before=src.read_bytes();second=self.folder/'second.docx';second.write_bytes(before)
  native=self.folder/'native.json';native.write_text(json.dumps({'sources':[{'id':'a','path':str(src)},{'id':'b','path':str(second)}],'make_master':True}),encoding='utf8')
  tasks=[{'id':'master','source':{'kind':'wps-job','path':str(native)},'outputs':['docx']},self.task('ready',self.document('ready.docx'))]
  first=run(self.rt,self.manifest(tasks));self.assertEqual(first['status'],'NEEDS_REVIEW',first)
  ready=first['tasks'][1]['files'];batch=self.folder/'output';home=batch/'tasks/master/wps';requests=json.loads((home/'review_requests.json').read_text('utf8'));responses=[]
  for request in requests['sources']:
   candidate=json.loads((home/'sources'/request['source_id']/'candidate.json').read_text('utf8'));decisions=[]
   facts={b['block_id']:b for b in json.loads((home/'sources'/request['source_id']/'scan/scan.json').read_text('utf8'))['blocks']}
   for row in candidate['blocks']:
    # The authored REF explanation is body text. The current conservative
    # engine may retain it as an unpaired caption; resolve this exact fixture
    # locally through the original source-bound decision protocol.
    reference_body=facts[row['block_id']].get('plain_text','').endswith('保留复杂域形式。')
    if row['status']!='review' and not reference_body:continue
    if not reference_body:self.assertEqual(blueprint.get(row['source_locator']['body_child_index']),'list_item')
    decisions.append(dict(block_id=row['block_id'],role='body' if reference_body else 'list_item',level=None,origin='model',decision_id='synthetic_'+request['source_id']+row['block_id'],reason='Authored list/reference fixture; synthetic review, no model call.',evidence_ids=[row['block_id']]))
   responses.append({**{k:request[k] for k in ('source_id','source_sha256','candidate_sha256')},'decisions':decisions})
  response=self.folder/'response.json';response.write_text(json.dumps({'sources':responses}),encoding='utf8');review(self.rt,batch,'master',response)
  done=resume(self.rt,batch);self.assertEqual(done['status'],'COMPLETE',done);self.assertEqual(done['tasks'][1]['files'],ready)
  self.assertEqual(len(done['tasks'][0]['deliverables']),3);self.assertEqual(src.read_bytes(),before);self.assertEqual(status(self.rt,batch)['integrity'],'PASS');self.assertFalse(list(batch.rglob('report.md')))
  state=json.loads((home/'state.json').read_text('utf8'));master=next(x for x in state['deliverables'] if x['id']=='master');package=Package(home/master['file']);facts=labels(package)
  self.assertIn('2.1.4.1.1.1',[f.get('label') for f in facts.values()]);seq=[f['cache'] for f in checked_fields(package.body) if f['instruction'].strip().startswith('SEQ 图')]
  self.assertEqual(seq,['1','2','3','4'])
 def test_one_failure_does_not_prevent_an_independent_good_task(self):
  bad=self.folder/'bad.docx';bad.write_bytes(b'not a docx');good=self.document()
  result=run(self.rt,self.manifest([self.task('bad',bad),self.task('good',good)]))
  self.assertEqual(result['status'],'PARTIAL');self.assertFalse(result['accepted'])
  self.assertEqual([t['status'] for t in result['tasks']],['FAIL','COMPLETE'])
  self.assertTrue(result['tasks'][1]['deliverables'])
 def test_review_pauses_one_task_and_resume_does_not_rebuild_completed_tasks(self):
  from hrs_bridge.batch import resume
  ambiguous=self.document('ambiguous.docx',True);good=self.document('good.docx')
  result=run(self.rt,self.manifest([self.task('review',ambiguous),self.task('good',good)]))
  self.assertEqual(result['status'],'NEEDS_REVIEW',result)
  self.assertEqual(result['tasks'][0]['status'],'NEEDS_REVIEW');self.assertEqual(result['tasks'][1]['status'],'COMPLETE')
  original=result['tasks'][1]['files'];resumed=resume(self.rt,self.folder/'output')
  self.assertEqual(resumed['tasks'][1]['files'],original);self.assertEqual(resumed['tasks'][0]['status'],'NEEDS_REVIEW')
 def test_changed_source_rejects_pending_resume_and_keeps_original_receipts(self):
  from hrs_bridge.batch import resume
  src=self.document(ambiguous=True);run(self.rt,self.manifest([self.task('review',src)]))
  src.write_bytes(src.read_bytes()+b'changed')
  result=resume(self.rt,self.folder/'output')
  self.assertEqual(result['tasks'][0]['status'],'FAIL');self.assertIn('changed',result['tasks'][0]['error'].lower())
  self.assertTrue((self.folder/'output/tasks/review/wps/review_requests.json').is_file())
 def test_changed_deliverable_is_detected_before_reuse(self):
  from hrs_bridge.batch import status
  result=run(self.rt,self.manifest([self.task('word',self.document())]))
  item=result['tasks'][0]['deliverables'][0];path=self.folder/'output'/item['path'];path.write_bytes(path.read_bytes()+b'tamper')
  checked=status(self.rt,self.folder/'output');self.assertEqual(checked['status'],'FAIL');self.assertEqual(checked['integrity'],'FAIL')
 def test_case_insensitive_duplicate_ids_are_rejected_before_start(self):
  src=self.document();job=self.manifest([self.task('A',src),self.task('a',src)])
  with self.assertRaisesRegex(ValueError,'Task IDs must be unique'):run(self.rt,job)
  self.assertFalse((self.folder/'output').exists())

 def test_source_bound_review_rejects_bad_response_then_accepts_a_correct_one(self):
  from hrs_bridge.batch import review,resume,status
  src=self.document(ambiguous=True);run(self.rt,self.manifest([self.task('review',src)]));batch=self.folder/'output'
  native=batch/'tasks/review/wps';requests=json.loads((native/'review_requests.json').read_text('utf8'))
  request=requests['sources'][0];candidate=json.loads((native/'sources/source_001/candidate.json').read_text('utf8'))
  scan=json.loads((native/'sources/source_001/scan/scan.json').read_text('utf8'))
  bid=next(r['block_id'] for r in scan['blocks'] if r.get('plain_text')=='1 FFT频谱分析')
  response={k:request[k] for k in ('source_id','source_sha256','candidate_sha256')}
  response['decisions']=[{'block_id':bid,'role':'heading','level':1,'origin':'model','decision_id':'synthetic-heading',
   'reason':'Controlled synthetic case: this text is the chapter title, not a measurement.','evidence_ids':[bid],'resolve_numeric_ambiguity':True}]
  path=self.folder/'response.json';invalid=json.loads(json.dumps(response));invalid['source_sha256']='0'*64
  path.write_text(json.dumps({'sources':[invalid]}),encoding='utf8')
  with self.assertRaises(ValueError):review(self.rt,batch,'review',path)
  self.assertEqual(status(self.rt,batch)['integrity'],'PASS')
  path.write_text(json.dumps({'sources':[response]}),encoding='utf8');review(self.rt,batch,'review',path)
  result=resume(self.rt,batch);self.assertEqual(result['status'],'COMPLETE',result)
  self.assertEqual(result['tasks'][0]['native_status']['receipt_integrity'],'PASS')

 def test_plain_markdown_relative_images_and_requested_outputs_share_existing_engines(self):
  from PIL import Image
  images=self.folder/'图片';images.mkdir();image=images/'测试 图.png';Image.new('RGB',(8,8),'blue').save(image)
  md=self.folder/'输入.md';text='# 技术报告\n\n## 方法\n\n正文值 1.2毫米。\n\n![说明](图片/测试%20图.png)\n'
  md.write_text(text,encoding='utf8');before=md.read_bytes()
  job=self.manifest([{'id':'md','source':{'kind':'markdown','path':str(md)},'outputs':['docx','hrs','html']}])
  result=run(self.rt,job);self.assertEqual(result['status'],'COMPLETE',result)
  task=result['tasks'][0];self.assertEqual(task['route'],'hrs-bridge');self.assertEqual(md.read_bytes(),before)
  self.assertEqual({x['kind'] for x in task['deliverables']},{'docx','hrs','html','provenance'})
  project=self.folder/'output/tasks/md/inputs/project'
  self.assertEqual(len(list((project/'assets').glob('*.png'))),1)

 def test_hrs_project_and_archive_html_are_explicit_input_routes(self):
  from hrs_bridge.browser import export
  project=self.folder/'project';(project/'assets').mkdir(parents=True);(project/'report.md').write_text('# 报告\n\n## 方法\n\n内容。\n',encoding='utf8')
  html=self.folder/'source.html';export(self.rt,project,html)
  tasks=[{'id':'project','source':{'kind':'hrs','path':str(project)},'outputs':['docx']},
   {'id':'archive','source':{'kind':'html','path':str(html)},'outputs':['docx','hrs']}]
  result=run(self.rt,self.manifest(tasks));self.assertEqual(result['status'],'COMPLETE',result)
  self.assertEqual([t['route'] for t in result['tasks']],['hrs-bridge','hrs-bridge'])

 def test_resume_refuses_simultaneous_writers_and_altered_task_specification(self):
  from hrs_bridge.batch import resume,status
  from project_commit import report_lock
  run(self.rt,self.manifest([self.task('review',self.document(ambiguous=True))]));folder=self.folder/'output'
  with report_lock(folder):
   with self.assertRaisesRegex(ValueError,'WRITE_BUSY'):resume(self.rt,folder)
  receipt=json.loads((folder/'WORKFLOW.json').read_text('utf8'));receipt['tasks'][0]['spec']['review_policy']='retain-unresolved'
  (folder/'WORKFLOW.json').write_text(json.dumps(receipt),encoding='utf8')
  with self.assertRaisesRegex(ValueError,'specification changed'):status(self.rt,folder)

 def test_bridge_word_import_freezes_its_sidecar_instead_of_losing_identity(self):
  from hrs_bridge.docx_export import convert
  project=self.folder/'origin';(project/'assets').mkdir(parents=True)
  original=b'# Report\n\n## Method\n\nContent.\n';(project/'report.md').write_bytes(original)
  source=self.folder/'own.docx';convert(self.rt,project,source)
  result=run(self.rt,self.manifest([{'id':'own','source':{'kind':'docx','path':str(source)},'outputs':['hrs']}]))
  self.assertEqual(result['status'],'COMPLETE',result)
  task=result['tasks'][0];back=self.folder/'output'/next(d['path'] for d in task['deliverables'] if d['kind']=='hrs')
  self.assertEqual(back.read_bytes(),original)
  self.assertTrue((self.folder/'output/tasks/own/inputs/source.docx.hrs-bridge.json').is_file())

 def test_external_word_to_hrs_can_resume_with_source_bound_structure_map(self):
  from hrs_bridge.batch import review,resume
  src=self.document(ambiguous=True)
  result=run(self.rt,self.manifest([{'id':'import','source':{'kind':'docx','path':str(src)},'outputs':['hrs']}]))
  self.assertEqual(result['status'],'NEEDS_REVIEW',result)
  record=json.loads(Path(result['tasks'][0]['review_requests']).read_text('utf8'))
  pending=record['review'][0];bid=pending['block']
  response={'schema':'hrs-word-structure-map-v1','source_sha256':record['source_sha256'],'roles':{
   bid:{'source_sha256':record['source_sha256'],'role':'heading','level':1,'origin':'model','reason':'Explicit synthetic chapter blueprint.',
        'decision_id':'synthetic-import','evidence_ids':[bid],'resolve_numeric_ambiguity':True}}}
  path=self.folder/'map.json';path.write_text(json.dumps(response),encoding='utf8')
  review(self.rt,self.folder/'output','import',path);finished=resume(self.rt,self.folder/'output')
  self.assertEqual(finished['status'],'COMPLETE',finished)

if __name__=='__main__':unittest.main()
