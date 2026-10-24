"""FB-01: data-table endings leave a template-sized gap to the next object."""
import json
from pathlib import Path
import tempfile
import unittest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from hrs_bridge.docx_export import convert
from hrs_bridge.docx_import import convert as restore
from hrs_bridge.browser import call
from test_html import runtime


class TableSpacing(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory();cls.root=Path(cls.tmp.name);cls.rt=runtime()
        src=cls.root/'source';(src/'assets').mkdir(parents=True)
        Image.new('RGB',(32,24),'navy').save(src/'assets/p.png')
        def table(n):return '```report-table\n'+json.dumps({'schema':'report-table-v1','id':'tbl-'+str(n),'caption':'表题'+str(n),'rows':[[{'text':t,'rowspan':1,'colspan':1} for t in row] for row in [['甲','乙'],['1','2']]]},ensure_ascii=False)+'\n```\n\n'
        group=call(cls.rt,'serialize',items=[{'type':'image-group','data':{'schema':'report-image-group-v1','id':'fig-group','caption':'组图','columns':2,'images':[{'src':'assets/p.png','caption':''}]*2,'showSubfigureLabels':False}}])[0]
        cls.raw=('# 表后间距\n\n'+table(1)+'表后正文\n\n'+table(2)+table(3)+group).encode('utf8')
        (src/'report.md').write_bytes(cls.raw);cls.src=src;cls.path=cls.root/'report.docx';convert(cls.rt,src,cls.path)

    @classmethod
    def tearDownClass(cls):cls.tmp.cleanup()

    def test_following_body_caption_and_group_have_spacing_without_changing_cell_text(self):
        """Body, next table caption and a borderless group all clear the preceding data grid."""
        d=Document(self.path);gap=d.styles['表题'].paragraph_format.space_before
        p=next(p for p in d.paragraphs if p.text=='表后正文')
        self.assertGreaterEqual(p.paragraph_format.space_before or 0,gap)
        caption=next(p for p in d.paragraphs if '表题3' in p.text)
        self.assertGreaterEqual(caption.paragraph_format.space_before or caption.style.paragraph_format.space_before or 0,gap)
        for cell in d.tables[-1].rows[0].cells:
            self.assertGreaterEqual(cell.paragraphs[0].paragraph_format.space_before or 0,gap)
        self.assertEqual(d.tables[0].cell(1,0).text,'1')

    def test_spacing_is_derived_layout_not_a_new_markdown_paragraph(self):
        """No spacer content enters the source or its exact same-source reverse conversion."""
        out=self.root/'back';restore(self.rt,self.path,out)
        self.assertEqual((out/'report.md').read_bytes(),self.raw)
        self.assertEqual((self.src/'report.md').read_bytes(),self.raw)


if __name__=='__main__':unittest.main()
