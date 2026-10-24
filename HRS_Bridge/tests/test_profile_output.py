"""PROFILE-01/WORD-01: literal source, template-owned numbering, abstract and code styles."""
from pathlib import Path
import json
import tempfile
import unittest
from docx import Document
from docx.oxml.ns import qn
from hrs_bridge.docx_import import escaped
from test_html import runtime


class LiteralWordText(unittest.TestCase):
    def test_word_literals_do_not_activate_references_or_math(self):
        """Text observed in Word remains text; real native objects use separate adapters."""
        rt=runtime()
        source=escaped('文字 [[参数]] @fig:fig-demo $E=mc^2$ ==重点==')
        self.assertEqual(rt.analyze((source+'\n').encode())['references'],[])
        from reference_context import masked_context
        self.assertIsNone(rt.contract.REF.search(masked_context(source)))
        self.assertIn(r'\$E\=mc^2\$',source)
        self.assertIn(r'\=\=重点\=\=',source)


class ProfileOutput(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from hrs_bridge.docx_export import convert
        cls.tmp=tempfile.TemporaryDirectory(prefix='bridge-profile-');cls.folder=Path(cls.tmp.name)
        cls.rt=runtime();cls.source=cls.folder/'source';cls.source.mkdir();(cls.source/'assets').mkdir()
        cls.raw='# 报告\n\n> 摘要：本次摘要。\n\n## 正文\n\n```python\nx = 1\n    print("@fig:fig-demo $E=mc^2$")\n```\n\n> 普通引用\n\n1. 列表\n'.encode('utf8')
        (cls.source/'report.md').write_bytes(cls.raw)
        (cls.source/'meta.json').write_text(json.dumps({'render_profile':'free'}),encoding='utf8')
        cls.path=cls.folder/'report.docx';cls.receipt=convert(cls.rt,cls.source,cls.path);cls.doc=Document(cls.path)

    @classmethod
    def tearDownClass(cls):cls.tmp.cleanup()

    def test_code_is_literal_and_ordinary_quote_uses_body_style(self):
        """Code stays editable text; ordinary quotes use template body.

        The current abstract policy (normal body with only its label bold) is
        exercised by WordOptions.test_02, rather than this obsolete style check.
        """
        code=next(p for p in self.doc.paragraphs if p.text.startswith('x = 1'))
        self.assertEqual(code.style.name,'HRS Code');self.assertIn('\n    print(',code.text)
        self.assertEqual(next(p for p in self.doc.paragraphs if p.text=='普通引用').style.name,'Normal')
        self.assertEqual((self.source/'report.md').read_bytes(),self.raw)

    def test_current_template_owns_styles_and_heading_numbering(self):
        """The default current mother template already has live numbering; export adds no outline."""
        original=Document(self.rt.wps/'templates/target_template.docx')
        for style in ['HRS Code','HRS Abstract']:self.assertIn(style,original.styles)
        nums=[original.styles['Heading '+str(i)].element.find('.//'+qn('w:numId')) for i in range(1,7)]
        self.assertTrue(all(n is not None for n in nums));self.assertEqual(len({n.get(qn('w:val')) for n in nums}),1)
        self.assertEqual([n.get(qn('w:val')) for n in nums],[self.doc.styles['Heading '+str(i)].element.find('.//'+qn('w:numId')).get(qn('w:val')) for i in range(1,7)])
        outlines=[n for n in self.doc.part.numbering_part.element.findall(qn('w:abstractNum')) if n.find(qn('w:name')) is not None and n.find(qn('w:name')).get(qn('w:val'))=='HRS Template Headings']
        self.assertEqual(len(outlines),1)

    def test_word_output_uses_engineering_rules_but_restores_source_profile(self):
        """A note's Word output is template-controlled; unchanged return keeps its original profile."""
        from hrs_bridge.docx_import import convert
        out=self.folder/'restored';convert(self.rt,self.path,out)
        self.assertEqual((out/'report.md').read_bytes(),self.raw)
        self.assertEqual(json.loads((out/'meta.json').read_text())['render_profile'],'free')


if __name__=='__main__':unittest.main()
