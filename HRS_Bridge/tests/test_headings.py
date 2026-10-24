"""HEAD-01/02: single-report title mapping and independent heading/list numbering."""
import json
from pathlib import Path
import tempfile
import unittest
from docx import Document
from docx.oxml.ns import qn
from hrs_bridge.docx_export import convert
from hrs_bridge.docx_import import convert as recover
from hrs_bridge.runtime import BridgeError, sha
from test_html import runtime


class HeadingNumbering(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix='bridge-heading-')
        cls.folder = Path(cls.temp.name)
        cls.project = cls.folder / 'source'
        cls.project.mkdir()
        (cls.project / 'assets').mkdir()
        cls.raw = ('# 单篇报告 {#sec-report}\n\n## 1 第一章 {#sec-first}\n\n### 1.1 第一节\n\n'
                   '#### 1.1.1 第三级\n\n##### 1.1.1.1 第四级\n\n'
                   '###### 1.1.1.1.1 第五级\n\n### 1.2 mm 测量范围\n\n## 第二章\n\n'
                   '### 末节\n').encode('utf8')
        (cls.project / 'report.md').write_bytes(cls.raw)
        cls.rt = runtime()
        cls.path = cls.folder / 'headings.docx'
        # This authored fixture deliberately contains manual heading prefixes;
        # the current HRS output gate requires explicit snapshot acceptance.
        cls.receipt = convert(cls.rt, cls.project, cls.path,accept_numbered_headings=True)
        cls.doc = Document(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_all_six_heading_styles_share_native_outline(self):
        """Applying an existing Heading style also supplies its correct live list level."""
        nums=[]
        for level in range(1,7):
            pr=self.doc.styles['Heading '+str(level)].element.find(qn('w:pPr'))
            num=pr.find(qn('w:numPr'))
            self.assertIsNotNone(num)
            self.assertEqual(num.find(qn('w:ilvl')).get(qn('w:val')),str(level-1))
            nums.append(num.find(qn('w:numId')).get(qn('w:val')))
        self.assertEqual(len(set(nums)),1)
        root=self.doc.part.numbering_part.element
        number=root.find(qn('w:num')+'[@'+qn('w:numId')+'="'+nums[0]+'"]')
        aid=number.find(qn('w:abstractNumId')).get(qn('w:val'))
        abstract=root.find(qn('w:abstractNum')+'[@'+qn('w:abstractNumId')+'="'+aid+'"]')
        self.assertEqual([e.find(qn('w:lvlText')).get(qn('w:val')) for e in abstract.findall(qn('w:lvl'))],
                         ['.'.join('%'+str(i) for i in range(1,n+1)) for n in range(1,7)])
        self.assertIsNone(self.doc.styles['Title'].element.find('.//'+qn('w:numPr')))

    def test_literal_chapter_prefix_does_not_duplicate_or_eat_quantities(self):
        """Only the exact expected chapter prefix is detached; numeric title content survives."""
        self.assertEqual([p.text for p in self.doc.paragraphs],
                         ['单篇报告','第一章','第一节','第三级','第四级','第五级','1.2 mm 测量范围','第二章','末节'])
        self.assertEqual((self.project/'report.md').read_bytes(),self.raw)

    def test_single_report_h1_is_title_and_h2_starts_chapters(self):
        """HEAD-02: the report title is unnumbered; Markdown H2 starts Word Heading 1."""
        self.assertEqual([p.style.name for p in self.doc.paragraphs],
                         ['Title','Heading 1','Heading 2','Heading 3','Heading 4','Heading 5',
                          'Heading 2','Heading 1','Heading 2'])

    def test_edited_chapter_returns_to_original_markdown_level_and_id(self):
        """HEAD-02: editing a Word Heading 1 must recover Markdown H2, not another title."""
        doc=Document(self.path)
        doc.paragraphs[1].text='更新章节'
        path=self.folder/'edited-heading.docx';doc.save(path)
        out=self.folder/'edited-heading'
        recover(self.rt,path,out,provenance=Path(self.receipt['provenance']))
        blocks=self.rt.analyze((out/'report.md').read_bytes(),out/'assets')['blocks']
        heading=next(b for b in blocks if b.get('id')=='sec-first')
        self.assertEqual((heading['level'],heading['text']),(2,'更新章节'))

    def test_edited_title_remains_markdown_h1(self):
        """HEAD-02: Word Title text edits retain the HRS title level and identity."""
        doc=Document(self.path);doc.paragraphs[0].text='更新报告题名'
        path=self.folder/'edited-title.docx';doc.save(path)
        out=self.folder/'edited-title'
        recover(self.rt,path,out,provenance=Path(self.receipt['provenance']))
        title=self.rt.analyze((out/'report.md').read_bytes(),out/'assets')['blocks'][0]
        self.assertEqual((title['level'],title['text'],title['id']),(1,'更新报告题名','sec-report'))

    def test_multiple_h1_requires_explicit_organization(self):
        """HEAD-02: multiple H1 reports cannot silently enter single-report title mapping."""
        project=self.folder/'multiple';project.mkdir();(project/'assets').mkdir()
        (project/'report.md').write_bytes('# 报告甲\n\n## 背景\n\n# 报告乙\n'.encode('utf8'))
        with self.assertRaisesRegex(BridgeError,'MULTIPLE_H1'):
            convert(self.rt,project,self.folder/'multiple.docx')

    def test_unchanged_native_numbered_document_recovers_source_bytes(self):
        """Derived Word outline labels must never overwrite the Markdown source snapshot."""
        out=self.folder/'restored'
        recover(self.rt,self.path,out)
        self.assertEqual((out/'report.md').read_bytes(),self.raw)

    def test_legacy_word_metadata_cannot_bypass_engineering_title_checks(self):
        """NUMBERING-02: obsolete unprojected multi-H1 metadata is not an exemption.

        Current native imports use source-bound overflow projection; its positive
        six-level roundtrip is covered by test_overflow_prefix.
        """
        project=self.folder/'word-origin';project.mkdir();(project/'assets').mkdir()
        raw=('# Word题名\n\n'+''.join('#'*i+' 原Word第'+str(i)+'级\n\n' for i in range(1,7))).encode('utf8')
        (project/'report.md').write_bytes(raw)
        title=self.rt.analyze(raw,project/'assets')['blocks'][0]
        origin={'schema':'hrs-word-origin-v1','markdown_sha256':sha(raw),'word_styles':[
            {'start':title['start'],'source_sha256':sha(title['source'].encode('utf8')),'style':'Title'}]}
        (project/'bridge-origin.json').write_text(json.dumps(origin),encoding='utf8')
        path=self.folder/'word-origin.docx'
        with self.assertRaisesRegex(BridgeError,'MULTIPLE_H1'):convert(self.rt,project,path)
        self.assertFalse(path.exists());self.assertEqual((project/'report.md').read_bytes(),raw)


MIXED_SOURCE = ('# 标题与列表专项\n\n## 摘要\n\n1. 列表甲\n2. 列表乙\n\n'
                '### 细节\n\n- 无序甲\n- 无序乙\n  1. 嵌套甲\n  2. 嵌套乙\n\n'
                '## 结论\n\n3. 后续甲\n4. 后续乙\n')


class HeadingAndLists(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory(prefix='bridge-heading-list-')
        cls.folder=Path(cls.temp.name);cls.project=cls.folder/'source'
        cls.project.mkdir();(cls.project/'assets').mkdir()
        (cls.project/'report.md').write_bytes(MIXED_SOURCE.encode('utf8'))
        cls.path=cls.folder/'mixed.docx'
        convert(runtime(),cls.project,cls.path)
        cls.doc=Document(cls.path)

    @classmethod
    def tearDownClass(cls):cls.temp.cleanup()

    def test_heading_and_list_numbers_are_disjoint(self):
        """HEAD-02: ordered, bullet and nested lists never share a heading counter."""
        heading_ids={self.doc.styles['Heading '+str(i)].element.find('.//'+qn('w:numId')).get(qn('w:val')) for i in range(1,7)}
        lists=[p for p in self.doc.paragraphs if p._p.find('.//'+qn('w:numPr')) is not None]
        self.assertEqual(len(lists),8)
        self.assertTrue(all(p.style.name=='Normal' for p in lists))
        self.assertTrue(all(p._p.find('.//'+qn('w:numId')).get(qn('w:val')) not in heading_ids for p in lists))
        self.assertEqual([p.style.name for p in self.doc.paragraphs if p.text in ('摘要','细节','结论')],
                         ['Heading 1','Heading 2','Heading 1'])

    def test_lists_keep_bullets_nesting_and_explicit_start(self):
        """HEAD-02: nested level and a separate list starting at 3 survive the mapping fix."""
        paragraphs={p.text:p for p in self.doc.paragraphs}
        nested=paragraphs['嵌套甲']._p.find('.//'+qn('w:ilvl'))
        self.assertEqual(nested.get(qn('w:val')),'1')
        number=paragraphs['后续甲']._p.find('.//'+qn('w:numId')).get(qn('w:val'))
        root=self.doc.part.numbering_part.element
        item=root.find(qn('w:num')+'[@'+qn('w:numId')+'="'+number+'"]')
        aid=item.find(qn('w:abstractNumId')).get(qn('w:val'))
        abstract=root.find(qn('w:abstractNum')+'[@'+qn('w:abstractNumId')+'="'+aid+'"]')
        self.assertEqual(abstract.find('.//'+qn('w:start')).get(qn('w:val')),'3')


if __name__=='__main__':unittest.main()
