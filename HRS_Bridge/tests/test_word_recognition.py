"""FB-03/04: reuse WPS candidates before projecting mixed Word structure."""
from pathlib import Path
import tempfile
import unittest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from PIL import Image
from hrs_bridge.docx_import import convert
from test_html import runtime


class WordRecognition(unittest.TestCase):
    def test_indented_subcaption_row_does_not_become_ordinary_body_baseline(self):
        """Spaces aligning two labels do not make a short caption count as long prose."""
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);rt=runtime();doc=Document();image=p/'image.png'
            Image.new('RGB',(60,40),'navy').save(image)
            # An explicit article title keeps this test focused on subcaption
            # layout, independently of the leading-title inference contract.
            doc.add_paragraph('双图试验报告',style='Title')
            row=doc.add_paragraph()
            for _ in range(2):row.add_run().add_picture(str(image),width=Inches(1))
            doc.add_paragraph('测试1'+' '*66+'测试2').paragraph_format.first_line_indent=Inches(.83)
            source=p/'source.docx';doc.save(source);convert(rt,source,p/'back')
            blocks=rt.analyze((p/'back/report.md').read_bytes(),p/'back/assets')['blocks']
            group=next(b for b in blocks if b['type']=='image-group')['data']
            self.assertEqual([v['caption'] for v in group['images']],['测试1','测试2'])

    def test_numbered_peer_and_unstyled_figure_caption_are_resolved_from_context(self):
        """A matching peer chapter stays a heading; centered text adjacent to an image is its caption."""
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);rt=runtime();doc=Document();doc.add_paragraph('题名',style='Title')
            doc.add_paragraph('1 模板结构',style='Heading 1');doc.add_paragraph('本节的正文内容。')
            image=p/'image.png';Image.new('RGB',(60,40),'navy').save(image)
            doc.add_paragraph().add_run().add_picture(str(image),width=Inches(1))
            doc.add_paragraph('结构示意',style='Normal').alignment=WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph('2. 第二章');doc.add_paragraph('本节包含后续的正文内容。')
            source=p/'source.docx';doc.save(source)
            receipt=convert(rt,source,p/'back');raw=(p/'back/report.md').read_bytes()
            blocks=rt.analyze(raw,p/'back/assets')['blocks']
            self.assertEqual([(b['level'],b['text']) for b in blocks if b['type']=='heading'],
                             [(1,'题名'),(2,'1 模板结构'),(2,'2. 第二章')])
            self.assertEqual(next(b['data']['caption'] for b in blocks if b['type']=='figure'),'结构示意')
            self.assertTrue((p/'back/word-recognition.json').is_file())
            self.assertEqual(receipt['recognition']['unresolved'],0)

    def test_same_row_images_and_aligned_subcaptions_do_not_need_manual_block_ids(self):
        """Two inline images in one row become one HRS group; aligned labels bind by column."""
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);rt=runtime();doc=Document();doc.add_paragraph('题名',style='Title')
            image=p/'image.png';Image.new('RGB',(60,40),'navy').save(image)
            row=doc.add_paragraph()
            for _ in range(2):row.add_run().add_picture(str(image),width=Inches(1))
            doc.add_paragraph('左侧说明\t右侧说明').alignment=WD_ALIGN_PARAGRAPH.CENTER
            source=p/'source.docx';doc.save(source)
            convert(rt,source,p/'back');blocks=rt.analyze((p/'back/report.md').read_bytes(),p/'back/assets')['blocks']
            group=next(b for b in blocks if b['type']=='image-group')['data']
            self.assertEqual((group['rows'],group['columns']),(1,2))
            self.assertEqual([v['caption'] for v in group['images']],['左侧说明','右侧说明'])
            self.assertFalse(group['showSubfigureLabels'])

    def test_real_heading_after_image_is_not_absorbed_as_caption(self):
        """Heading style is structural evidence; centered alignment alone cannot override it."""
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);rt=runtime();doc=Document()
            image=p/'image.png';Image.new('RGB',(60,40),'navy').save(image)
            doc.add_paragraph().add_run().add_picture(str(image),width=Inches(1))
            doc.add_paragraph('真实章节',style='Heading 2').alignment=WD_ALIGN_PARAGRAPH.CENTER
            source=p/'source.docx';doc.save(source)
            convert(rt,source,p/'back');blocks=rt.analyze((p/'back/report.md').read_bytes(),p/'back/assets')['blocks']
            self.assertEqual(next(b['data']['caption'] for b in blocks if b['type']=='figure'),'')
            self.assertEqual(next(b['level'] for b in blocks if b['type']=='heading'),3)

    def test_list_introduction_prevents_typed_steps_becoming_peer_chapters(self):
        """Chapter numbering does not turn the following explicitly introduced list into headings."""
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);rt=runtime();doc=Document();doc.add_paragraph('1 概述',style='Heading 1')
            doc.add_paragraph('下面操作步骤如下：');doc.add_paragraph('2. 打开软件');doc.add_paragraph('3. 检查结果')
            source=p/'source.docx';doc.save(source)
            convert(rt,source,p/'back');blocks=rt.analyze((p/'back/report.md').read_bytes(),p/'back/assets')['blocks']
            self.assertEqual(len([b for b in blocks if b['type']=='heading']),1)


if __name__=='__main__':unittest.main()
