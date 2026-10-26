"""WORD-HYPERLINK-01: observed DOCX link boundaries, source and target fidelity."""
from pathlib import Path
import json,tempfile,unittest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from hrs_bridge.docx_import import convert,NativeReader
from hrs_bridge.runtime import BridgeError
from hrs_bridge.ooxml import semantic_hash
from test_html import runtime


def link(p,label,target=None,anchor=None,bold=False):
    node=OxmlElement('w:hyperlink')
    if target is not None:node.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
    if anchor is not None:node.set(qn('w:anchor'),anchor)
    for text in label if isinstance(label,list) else [label]:
        r=OxmlElement('w:r')
        if bold:
            prop=OxmlElement('w:rPr');prop.append(OxmlElement('w:b'));r.append(prop)
        t=OxmlElement('w:t');t.text=text;r.append(t);node.append(r)
    p._p.append(node);return node


def bookmark(p,name,number=50):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(number));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number));p._p.insert(0,start);p._p.append(end)


class NativeHyperlinks(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='bridge-links-');self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.rt=runtime();self.doc=Document();self.doc.add_paragraph('链接试验',style='Title')

    def run_import(self):
        source=self.root/'输入.docx';self.doc.save(source);before=source.read_bytes()
        result=convert(self.rt,source,self.root/'project')
        self.assertEqual(source.read_bytes(),before)
        text=(self.root/'project/report.md').read_text('utf8')
        return result,text,self.rt.analyze(text.encode('utf8'),self.root/'project/assets')

    def test_external_links_keep_multiple_runs_format_and_adjacent_literal_text(self):
        p=self.doc.add_paragraph('前文：');link(p,['链接','😀'],'https://example.org/a?x=1&y=2',bold=True)
        p.add_run('，中间原文；');link(p,'第二个','https://example.org/含 空格');p.add_run('。')
        result,text,_=self.run_import()
        self.assertIn('[**链接😀**](<https://example.org/a?x=1&y=2>)',text)
        self.assertIn('[第二个](<https://example.org/%E5%90%AB%20%E7%A9%BA%E6%A0%BC>)',text)
        self.assertIn('，中间原文；',text);self.assertEqual(result['status'],'COMPLETE')

    def test_forward_heading_bookmark_maps_to_stable_target(self):
        p=self.doc.add_paragraph('请参看');link(p,'跳至目的',anchor='_TocPurpose');p.add_run('。')
        p=self.doc.add_heading('目的',level=1);bookmark(p,'_TocPurpose')
        result,text,ir=self.run_import();heading=next(b for b in ir['blocks'] if b['type']=='heading' and b['level']==2)
        self.assertIn('[跳至目的](<#object-'+heading['id']+'>)',text)
        self.assertEqual(result['status'],'COMPLETE')

    def test_table_and_whole_caption_link_keep_target_without_old_caption_number(self):
        link(self.doc.add_paragraph(),'表1 参数表',target='https://example.org/table')
        table=self.doc.add_table(rows=1,cols=1);link(table.cell(0,0).paragraphs[0],'资料入口',target='mailto:sample@example.org')
        result,text,ir=self.run_import();data=next(b['data'] for b in ir['blocks'] if b['type']=='table')
        self.assertEqual(data['caption'],'[参数表](<https://example.org/table>)')
        self.assertEqual(data['rows'][0][0]['text'],'[资料入口](<mailto:sample@example.org>)')
        self.assertEqual(result['status'],'COMPLETE')

    def test_unknown_internal_target_is_visible_and_reviewable_not_a_broken_link(self):
        link(self.doc.add_paragraph(),'未定位的原文字',anchor='_MissingTarget')
        result,text,_=self.run_import();self.assertIn('未定位的原文字',text)
        self.assertNotIn('](<#',text);self.assertEqual(result['status'],'NEEDS_REVIEW')
        diagnostics=[d for d in result['diagnostics'] if d['code']=='WORD_HYPERLINK_UNRESOLVED']
        self.assertEqual(diagnostics[0]['anchor'],'_MissingTarget')
        record=json.loads((self.root/'project/word-recognition.json').read_text('utf8'))
        self.assertTrue(any(v.get('role')=='hyperlink' for v in record['review']))

    def test_unsafe_and_relative_targets_are_not_activated_or_silently_discarded(self):
        p=self.doc.add_paragraph();link(p,'危险入口',target='javascript:alert(1)');p.add_run(' 原文 ');link(p,'相对原链接',target='../docs/paper.pdf')
        result,text,_=self.run_import()
        self.assertIn('危险入口 原文 相对原链接',text);self.assertNotIn('](',text)
        self.assertEqual(result['status'],'NEEDS_REVIEW')
        self.assertEqual({d.get('target') for d in result['diagnostics'] if d['code']=='WORD_HYPERLINK_UNRESOLVED'}, {'javascript:alert(1)','../docs/paper.pdf'})

    def test_semantic_fingerprint_detects_native_internal_destination_change(self):
        p=self.doc.add_paragraph();node=link(p,'相同显示文字',anchor='First')
        before=semantic_hash([p._p],{},{});node.set(qn('w:anchor'),'Second')
        self.assertNotEqual(before,semantic_hash([p._p],{},{}))

    def test_scanner_range_audit_detects_tampering_and_keeps_number_fact_audit(self):
        from wps_pipeline.verify_scan_fidelity import verify
        p=self.doc.add_paragraph('前文');link(p,['两段','文字'],'https://example.org');p.add_run('后文。')
        source=self.root/'range.docx';self.doc.save(source);assets=self.root/'assets';assets.mkdir()
        reader=NativeReader(self.rt,source,self.root/'work',assets)
        passed=verify(source,self.root/'work/scan');self.assertEqual(passed['status'],'PASS',passed)
        self.assertIsNotNone(passed['recognition_fact_audit'])
        path=self.root/'work/scan/scan.json';scan=json.loads(path.read_text('utf8'))
        block=next(b for b in scan['blocks'] if b.get('hyperlink_ranges'))
        block['hyperlink_ranges'][0]['end']-=1;path.write_text(json.dumps(scan),encoding='utf8')
        rejected=verify(source,self.root/'work/scan');self.assertEqual(rejected['status'],'FAIL')
        self.assertTrue(any('hyperlink_range_label' in v for v in rejected['errors']))

    def test_link_range_cannot_overlap_or_select_an_unrelated_segment(self):
        p=self.doc.add_paragraph('前文');link(p,'链接','https://example.org');p.add_run('后文。')
        source=self.root/'bounds.docx';self.doc.save(source);assets=self.root/'assets';assets.mkdir()
        reader=NativeReader(self.rt,source,self.root/'work',assets)
        block=next(b for b in reader.top if b.get('hyperlink_ranges'))
        block['hyperlink_ranges'][0]['start']=-1
        with self.assertRaisesRegex(BridgeError,'WORD_LINK_RANGE_INVALID'):reader.paragraph(block)

    def test_ignored_proof_marker_does_not_shift_link_range(self):
        p=self.doc.add_paragraph('前文');node=link(p,'链接','https://example.org')
        proof=OxmlElement('w:proofErr');proof.set(qn('w:type'),'spellStart');node.insert(0,proof);p.add_run('后文。')
        result,text,_=self.run_import();self.assertIn('前文[链接](<https://example.org>)后文。',text)
        self.assertEqual(result['status'],'COMPLETE')

    def test_external_document_location_is_part_of_the_target(self):
        p=self.doc.add_paragraph('参见');node=link(p,'定位资料','https://example.org/report')
        node.set(qn('w:docLocation'),'Part 2');p.add_run('。')
        result,text,_=self.run_import();self.assertIn('[定位资料](<https://example.org/report#Part%202>)',text)
        self.assertEqual(result['status'],'COMPLETE')


if __name__=='__main__':unittest.main()
