"""FB-02: one HRS title, shifted chapters and list projection beyond H6."""
from pathlib import Path
import json
import tempfile
import unittest
from docx import Document
from hrs_bridge.docx_import import convert as read_word
from hrs_bridge.docx_export import convert as write_word
from test_html import runtime, ROOT


class HeadingProjection(unittest.TestCase):
    def test_title_and_body_levels_do_not_compete_for_h1(self):
        """Word Title stays H1; Word body levels 1..5 become H2..6, level 6 a list."""
        with tempfile.TemporaryDirectory() as temp:
            folder=Path(temp);rt=runtime()
            doc=Document(ROOT/'fixtures/target_template.docx')
            for p in list(doc.paragraphs):p._element.getparent().remove(p._element)
            doc.add_paragraph('报告题名',style='Title')
            for level in range(1,7):doc.add_paragraph('原正文第'+str(level)+'级',style='Heading '+str(level))
            doc.add_paragraph('最深标题之后的正文。')
            source=folder/'source.docx';doc.save(source)
            read_word(rt,source,folder/'back')
            raw=(folder/'back/report.md').read_bytes()
            ir=rt.analyze(raw,folder/'back/assets')
            heads=[b for b in ir['blocks'] if b['type']=='heading']
            self.assertEqual([b['level'] for b in heads],[1,2,3,4,5,6])
            self.assertIn('- 原正文第6级\n',raw.decode('utf8'))
            origin=json.loads((folder/'back/bridge-origin.json').read_text('utf8'))
            self.assertEqual([b['style'] for b in origin['word_styles']],['Title']+['Heading '+str(i) for i in range(1,7)])
            write_word(rt,folder/'back',folder/'again.docx')
            paragraphs=Document(folder/'again.docx').paragraphs
            self.assertEqual([p.style.name for p in paragraphs[:7]],['Title']+['Heading '+str(i) for i in range(1,7)])
            self.assertEqual(paragraphs[-1].text,'最深标题之后的正文。')

    def test_report_without_explicit_title_does_not_lose_a_chapter(self):
        """A missing Title never turns the first body chapter into the report title."""
        with tempfile.TemporaryDirectory() as temp:
            folder=Path(temp);rt=runtime();doc=Document()
            doc.add_paragraph('正文第一章',style='Heading 1')
            source=folder/'source.docx';doc.save(source)
            read_word(rt,source,folder/'back')
            self.assertIn('## 正文第一章',(folder/'back/report.md').read_text('utf8'))


if __name__=='__main__':unittest.main()
